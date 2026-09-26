"""Enterprise recognition tests use synthetic files only (no customer data)."""

from __future__ import annotations

import io
from types import SimpleNamespace
from typing import Any

from openpyxl import Workbook

from backend.application.import_recognition import decimal_text
from backend.application.reference_transfer import validate_transfer
from backend.application.subsidiary_report import build_rows
from backend.infrastructure.excel.reference_workbook import read_reference


def sample(*, week: bool = True, hidden: bool = False) -> dict[str, Any]:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "Отчёт 2026"
    for address, text in {
        "B6": "№ п/п",
        "C6": "Обозначение",
        "D6": "Наименование",
        "E6": "Входимость",
        "F6": "Производитель",
        "H6": "Объём поставок по договору",
        "CU6": "Факт",
    }.items():
        sheet[address] = text
    sheet["CU2"] = "Сентябрь 2026"
    if week:
        sheet["CU3"] = "1 неделя"
    sheet["C4"] = "=1/2"
    sheet["C4"].number_format = "0.0%"
    for address, value in {
        "B8": "01.1",
        "C8": "000.010",
        "D8": "Деталь",
        "E8": 2,
        "F8": "Первый",
        "H8": 30,
        "CU8": "=80+50",
        "F9": "Второй",
        "H9": 20,
        "CU9": 0,
    }.items():
        sheet[address] = str(value) if isinstance(value, str) else int(value)  # type: ignore[call-overload]
    for col in "BCDE":
        sheet.merge_cells(f"{col}8:{col}9")
    sheet.row_dimensions[9].hidden = True
    sheet.column_dimensions["CU"].hidden = True
    if hidden:
        extra = book.create_sheet("Дополнение")
        extra.sheet_state = "hidden"
        extra["A1"] = "Пояснения"
    stream = io.BytesIO()
    book.save(stream)
    return read_reference(stream.getvalue())


def matrix() -> dict[str, Any]:
    coordinate = {
        "report_type": "SUBSIDIARY",
        "organization_id": "1",
        "component_id": "1",
        "metric_code": "SUB_SUPPLIED_PRIMARY",
        "period_start": "2026-09-01",
    }
    return {
        "report_type": "SUBSIDIARY",
        "rows": [
            {
                "supplier_id": "PRIMARY",
                "left_values": {"designation": "000.010", "position": "Деталь", "party": "Первый"},
                "cells": [
                    {
                        "coordinate": coordinate,
                        "state": {"access": "editable"},
                        "value": {"kind": "DATA_NOT_PROVIDED"},
                    }
                ],
            }
        ],
        "time_columns": [{"id": "2026-09-01-SUPPLIED", "period_end": "2026-09-06"}],
    }


def test_full_extent_roles_formulas_merged_suppliers_hidden_data() -> None:
    doc = sample(hidden=True)
    assert doc["sheets"][0]["columns"] == 99
    assert doc["sheets"][0]["hidden_rows"] == [9]
    assert doc["sheets"][0]["hidden_columns"] == ["CU"]
    rec = doc["recognition"]
    assert rec["fields"]["0:C4"]["role"] == "CONTROL"
    assert rec["fields"]["0:C4"]["unit"] == "%"
    assert "0:C4" not in rec["sources"]
    assert rec["sources"]["0:CU8"]["value"] == "130"
    assert rec["sources"]["0:CU8"]["formula_status"] == "CALCULATED_NO_CACHE"
    assert rec["sources"]["0:CU9"]["value"] == "0"
    assert rec["sources"]["0:CU8"]["period"] is None
    action = rec["structural_actions"][0]
    assert action["position"]["code"] == "000.010"
    assert action["weekly_supply"]
    assert [s["name"] for s in action["suppliers"]] == ["Первый", "Второй"]
    assert any(i["code"] == "SHEET_REQUIRED" for i in rec["issues"])


def test_empty_mapping_is_russian_and_count_has_no_extra_empty_error() -> None:
    doc = sample()
    checked = validate_transfer(doc, matrix(), [{"source": "0:CU8", "coordinate": {}}])
    assert any(i["message"] == "Выберите позицию и показатель" for i in checked["issues"])
    assert not any("unsupported report_type" in i["message"] for i in checked["issues"])
    assert not any(i["code"] == "EMPTY_TRANSFER" for i in checked["issues"])


def test_manual_mapping_cannot_ignore_exact_end_or_reverse_supply() -> None:
    doc = sample()
    target = matrix()
    coord = target["rows"][0]["cells"][0]["coordinate"]
    mappings = [
        {"source": "0:CU8", "coordinate": coord, "quantity": "130", "confirmed": True},
        {"source": "0:CU9", "skip_reason": "Отдельный изготовитель"},
    ]
    rules = {"0:CU": {"start": "2026-09-01", "end": "2026-09-07", "reason": "Исходный календарь"}}
    checked = validate_transfer(doc, target, mappings, rules)
    assert checked["issues"][0]["code"] == "PERIOD_REQUIRED"
    rules["0:CU"]["end"] = "2026-09-06"
    checked = validate_transfer(doc, target, mappings, rules)
    assert not checked["issues"]
    assert checked["changes"][0]["value"]["quantity"] == "130"
    coord["metric_code"] = "SUB_USED_PRIMARY"
    checked = validate_transfer(doc, target, mappings, rules)
    assert "Поставку нельзя" in checked["issues"][0]["message"]


def test_numeric_strings_keep_blank_zero_and_locale_distinct() -> None:
    assert decimal_text("1\u00a0234,50") == "1234.50"
    assert decimal_text("0") == "0"
    import pytest

    for raw in ("", "—", "2026-09-01", "000.001.2"):
        with pytest.raises(ValueError):
            decimal_text(raw)


def test_weekly_supply_increases_balance_and_future_blanks_stay_unknown() -> None:
    group = SimpleNamespace(id=1, party_name="Первый", position_name="Деталь")
    detail = {
        "number": "1",
        "designation": "000.010",
        "weekly_supply": True,
        "suppliers": [{"id": "PRIMARY", "name": "Первый", "contract": "", "archived": False}],
    }
    config = {"subsidiary": detail, "category": "UNSPECIFIED", "image": "", "norm": "1"}
    values = {
        ("SUB_OPENING", "2026-09-01"): "100",
        ("SUB_SUPPLIED_PRIMARY", "2026-09-01"): "20",
        ("SUB_USED_PRIMARY", "2026-09-01"): "5",
    }

    def read(g: Any, metric: str, day: str, column: str, editable: bool) -> dict[str, Any]:
        value = values.get((metric, day))
        return {
            "coordinate": {"metric_code": metric, "period_start": day},
            "time_column_id": column,
            "state": {"access": "editable" if editable else "calculated"},
            "value": {"kind": "DATA_NOT_PROVIDED"}
            if value is None
            else {"kind": "QUANTITY", "quantity": value},
        }

    result = build_rows([(group, config)], 2026, read, {})
    balances = result["rows"][0]["stock_by_week"]
    assert balances["2026-09-01"]["quantity"] == "115"
    assert balances["2026-09-07"]["kind"] == "DATA_NOT_PROVIDED"
    assert any(c["kind"] == "SUPPLIED" for c in result["time_columns"])


def test_explicit_structure_correction_updates_sources_without_mutating_original() -> None:
    from backend.application.import_recognition import apply_structure_overrides

    rec = sample()["recognition"]
    action = rec["structural_actions"][0]
    action["errors"] = [{"code": "MANUFACTURER_REQUIRED", "message": "Укажите изготовителя"}]
    corrected = apply_structure_overrides(
        rec,
        {
            action["source_key"]: {
                "manufacturer": "Подтверждённый изготовитель",
                "reason": "Уточнение владельца отчёта",
            }
        },
    )
    assert (
        corrected["sources"]["0:CU8"]["position"]["manufacturer"] == "Подтверждённый изготовитель"
    )
    assert corrected["structural_actions"][0]["errors"] == []
    assert rec["sources"]["0:CU8"]["position"]["manufacturer"] == "Первый"


def test_russian_march_is_not_may_and_iso_date_is_not_day_range() -> None:
    from backend.application.import_recognition import parse_period, profile

    march, _ = parse_period(["Март 2026"], None, profile())
    may, _ = parse_period(["Май 2026"], None, profile())
    august, _ = parse_period(["2025-08-01"], None, profile())
    assert march is not None and march["start"] == "2026-03-01"
    assert may is not None and may["start"] == "2026-05-01"
    assert august is not None and august["start"] == "2025-08-01"
