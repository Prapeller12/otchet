from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import Workbook, load_workbook
from openpyxl.worksheet.worksheet import Worksheet

from backend.api.working_reference_bridge import WorkingReferenceApplicationBridge
from backend.infrastructure.database.sqlite_reference_reports import ReferenceReports
from backend.infrastructure.excel.reference_workbook import read_reference

ROOT = Path(__file__).resolve().parents[3]


def fixture_book(path: Path, kind: str) -> None:
    wb = Workbook()
    sheet = cast(Worksheet, wb.active)
    sheet["E7"] = "Входимость в изделие"
    sheet["J7"] = "План" if kind == "HEAD_SITE" else "Общий дефицит"
    sheet["K7"] = "Факт"
    sheet["A1"] = "План 2026"
    sheet["J1"] = "Январь 2025"
    sheet["C9"] = "000.001"
    sheet["F9"] = "Изготовитель"
    sheet["L9"] = 375
    sheet["M9"] = 300
    sheet["N9"] = 0
    sheet["I9"] = "=SUM(L9:N9)"
    sheet.merge_cells("C9:C10")
    sheet["F10"] = "Другой изготовитель"
    wb.save(path)
    wb.close()


def bridge(tmp_path: Path) -> WorkingReferenceApplicationBridge:
    return WorkingReferenceApplicationBridge(
        tmp_path / "data/app.db",
        migrations_directory=ROOT / "backend/migrations",
        definitions_directory=ROOT / "resources/report-definitions",
    )


def unwrap(envelope: dict[str, Any]) -> Any:
    assert envelope["ok"], envelope
    return envelope["data"]


@pytest.mark.parametrize("kind", ["HEAD_SITE", "SUBSIDIARY"])
def test_reference_import_edit_restart_export(tmp_path: Path, kind: str) -> None:
    source = tmp_path / "source.XLSX"
    fixture_book(source, kind)
    app = bridge(tmp_path)
    app.configure_excel_dialogs(
        open_file=lambda: source, save_file=lambda _: tmp_path / "export.xlsx"
    )
    query = {"report_type": "DAILY_MOVEMENT", "organization_id": "1"}
    preview = unwrap(app.validate_import(query))
    assert preview["reference_workbook"]["report_type"] == kind
    assert unwrap(app.reference_report({"action": "list", "organization_id": "1"})) == []
    identity = unwrap(app.commit_import({"batch_id": preview["batch_id"]}))["reference_workbook_id"]
    repository = ReferenceReports(tmp_path / "data/app.db")
    loaded = repository.get(identity, 1)
    assert loaded["sheets"][0]["cells"]["I9"]["display"] == "675"
    assert loaded["sheets"][0]["cells"]["N9"]["display"] == "0"
    assert "O9" not in loaded["sheets"][0]["cells"]
    with pytest.raises(ValueError):
        repository.get(identity, 2)
    assert unwrap(app.validate_import(query))["already_imported"]
    assert len(repository.list_reports(1)) == 1
    for address in ("I9", "J1", "C10"):
        with pytest.raises(ValueError):
            repository.save(identity, 1, 0, [{"sheet": 0, "address": address, "value": "1"}])
    changed = repository.save(identity, 1, 0, [{"sheet": 0, "address": "L9", "value": "400"}])
    assert changed["sheets"][0]["cells"]["I9"]["display"] == "700"
    with pytest.raises(ValueError):
        repository.save(identity, 1, 0, [{"sheet": 0, "address": "M9", "value": "1"}])
    fresh = bridge(tmp_path)
    loaded = unwrap(
        fresh.reference_report({"action": "get", "organization_id": "1", "id": identity})
    )
    assert loaded["revision"] == 1
    repository.export(identity, 1, tmp_path / "export.xlsx")
    wb = load_workbook(tmp_path / "export.xlsx")
    assert cast(Worksheet, wb.active)["L9"].value == 400
    assert cast(Worksheet, wb.active)["I9"].value == "=SUM(L9:N9)"
    assert "C9:C10" in cast(Worksheet, wb.active).merged_cells
    wb.close()


def test_rejects_unknown_formulas_without_commit(tmp_path: Path) -> None:
    path = tmp_path / "bad.xlsx"
    fixture_book(path, "SUBSIDIARY")
    wb = load_workbook(path)
    cast(Worksheet, wb.active)["I9"] = '=WEBSERVICE("https://example.com")'
    wb.save(path)
    wb.close()
    assert read_reference(path.read_bytes())["errors"]
    app = bridge(tmp_path)
    app.configure_excel_dialogs(open_file=lambda: path, save_file=lambda _: None)
    assert not app.validate_import({"report_type": "SUBSIDIARY", "organization_id": "1"})["ok"]
    assert unwrap(app.reference_report({"action": "list", "organization_id": "1"})) == []
