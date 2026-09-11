from __future__ import annotations

from pathlib import Path
from typing import Any, cast
from uuid import uuid4

import pytest

from backend.api.working_reference_bridge import WorkingReferenceApplicationBridge
from backend.application.monthly_report import monthly_snapshot
from backend.infrastructure.database.migrator import apply_migrations, connect_sqlite
from backend.infrastructure.excel.openpyxl_matrix_workbook import OpenpyxlMatrixWorkbookAdapter
from backend.infrastructure.monthly_pdf import render_monthly_pdf

ROOT = Path(__file__).resolve().parents[3]
QUERY = {"report_type": "SUBSIDIARY", "organization_id": "1", "year": 2026}


def data(response: dict[str, Any]) -> Any:
    assert response["ok"], response
    return response["data"]


def app_at(path: Path) -> WorkingReferenceApplicationBridge:
    with connect_sqlite(path / "reporting.db") as connection:
        apply_migrations(connection, ROOT / "backend/migrations")
    return WorkingReferenceApplicationBridge(
        path / "reporting.db",
        migrations_directory=ROOT / "backend/migrations",
        definitions_directory=ROOT / "resources/report-definitions",
        backups_directory=path / "backups",
        inbox_directory=path / "inbox",
    )


def setup(app: WorkingReferenceApplicationBridge) -> dict[str, Any]:
    layout = data(app.get_report_layout({k: v for k, v in QUERY.items() if k != "year"}))
    row = layout["rows"][0]
    row["position_name"] = "Проверочная деталь"
    row["configuration"]["norm"] = "4"
    row["configuration"]["subsidiary"] = {
        "number": "1.1",
        "designation": "000.01",
        "suppliers": [
            {"id": "A", "name": "Производитель А", "contract": "2500", "archived": False},
            {"id": "B", "name": "Производитель Б", "contract": "2000", "archived": False},
        ],
    }
    data(
        app.save_report_layout({"report_type": "SUBSIDIARY", "organization_id": "1", "rows": [row]})
    )
    matrix = data(app.get_report_matrix(QUERY))
    data(
        app.save_report_presentation(
            {
                "report_type": "SUBSIDIARY",
                "organization_id": "1",
                "expected_revision": matrix["matrix_revision"],
                "plans": {"2026-09": "1000"},
            }
        )
    )
    return cast(dict[str, Any], data(app.get_report_matrix(QUERY)))


def write(
    app: WorkingReferenceApplicationBridge,
    matrix: dict[str, Any],
    entries: list[tuple[int, str, str]],
) -> dict[str, Any]:
    indices = {c["id"]: i for i, c in enumerate(matrix["time_columns"])}
    changes = [
        {
            "coordinate": matrix["rows"][row]["cells"][indices[column]]["coordinate"],
            "value": {"kind": "QUANTITY", "quantity": value},
        }
        for row, column, value in entries
    ]
    preview = data(app.get_report_matrix({**QUERY, "preview_changes": changes}))
    data(
        app.save_report_cells(
            {
                **QUERY,
                "base_revision": matrix["matrix_revision"],
                "idempotency_key": uuid4().hex,
                "changes": changes,
            }
        )
    )
    result = data(app.get_report_matrix(QUERY))
    assert preview["rows"] == result["rows"]
    return cast(dict[str, Any], result)


def value(matrix: dict[str, Any], column: str, row: int = 0) -> str | None:
    return cast(
        str | None,
        next(c for c in matrix["rows"][row]["cells"] if c["column_id"] == column)["value"].get(
            "quantity"
        ),
    )


def test_multiple_suppliers_consumption_and_monthly_variance(tmp_path: Path) -> None:
    app = app_at(tmp_path)
    matrix = setup(app)
    assert value(matrix, "2026-09-VARIANCE") is None
    matrix = write(
        app,
        matrix,
        [
            (0, "2026-09-OPENING", "500"),
            (0, "2026-09-RECEIVED", "2000"),
            (1, "2026-09-RECEIVED", "1000"),
            (0, "2026-09-01", "100"),
            (1, "2026-09-07", "200"),
            (0, "2026-09-28", "50"),
        ],
    )
    assert value(matrix, "2026-09-STOCK") == "3150"
    assert value(matrix, "2026-09-VARIANCE") == "-500"
    assert matrix["rows"][0]["stock_by_week"]["2026-09-01"]["quantity"] == "3400"
    assert matrix["rows"][0]["stock_by_week"]["2026-09-07"]["quantity"] == "3200"
    assert (
        len(
            [
                c
                for c in matrix["time_columns"]
                if c["group_label"] == "2026-09" and c["kind"] == "USED"
            ]
        )
        == 5
    )
    assert data(app_at(tmp_path).get_report_matrix(QUERY))["rows"] == matrix["rows"]
    # Consumption changes stock but must not subtract from monthly variance twice.
    matrix = write(app, matrix, [(0, "2026-09-28", "1000")])
    assert value(matrix, "2026-09-STOCK") == "2200"
    assert value(matrix, "2026-09-VARIANCE") == "-500"
    assert value(matrix, "2026-10-STOCK") is None
    workbook = tmp_path / "roundtrip.xlsx"
    adapter = OpenpyxlMatrixWorkbookAdapter()
    assert adapter.write(workbook, matrix) > 0
    parsed = adapter.parse(workbook, report_type="SUBSIDIARY", organization_id=1, matrix=matrix)
    assert not parsed.issues
    assert any(c.value.quantity == "2000" for c in parsed.cells)
    snapshot = monthly_snapshot(matrix, 9, "Общество")
    assert len([c for c in snapshot["columns"] if c["kind"] == "USED"]) == 5
    selected = monthly_snapshot(matrix, 9, "Общество", "2026-09-07")
    assert selected["as_of"] == "2026-09-13"
    assert selected["rows"][0]["values"][2] == "3200"
    with pytest.raises(ValueError):
        monthly_snapshot(matrix, 9, "Общество", "2026-08-01")
    status = data(app.get_report_verification({**QUERY, "month": 9, "week_start": "2026-09-07"}))
    data(
        app.verify_report(
            {
                **QUERY,
                "month": 9,
                "week_start": "2026-09-07",
                "signer_name": "Проверка",
                "confirmed": True,
                "snapshot_sha256": status["snapshot_sha256"],
            }
        )
    )
    assert (
        data(app.get_report_verification({**QUERY, "month": 9, "week_start": "2026-09-28"}))[
            "status"
        ]
        == "STALE"
    )
    font = ROOT / "resources/fonts/DejaVuSerif.ttf"
    if not font.exists():
        font = next((ROOT / "resources").rglob("*.ttf"))
    pdf = render_monthly_pdf(snapshot, {"status": "UNVERIFIED"}, font)
    assert pdf.startswith(b"%PDF-")
    (tmp_path / "subsidiary.pdf").write_bytes(pdf)


def test_zero_receipts_archive_and_plan_conflict(tmp_path: Path) -> None:
    app = app_at(tmp_path)
    matrix = setup(app)
    matrix = write(app, matrix, [(0, "2026-09-OPENING", "0"), (0, "2026-09-RECEIVED", "0")])
    assert value(matrix, "2026-09-VARIANCE") is None  # supplier B not reported
    matrix = write(app, matrix, [(1, "2026-09-RECEIVED", "0")])
    assert value(matrix, "2026-09-VARIANCE") == "-4000"
    layout = data(app.get_report_layout({"report_type": "SUBSIDIARY", "organization_id": "1"}))
    layout["rows"][0]["configuration"]["subsidiary"]["suppliers"][1]["archived"] = True
    data(
        app.save_report_layout(
            {"report_type": "SUBSIDIARY", "organization_id": "1", "rows": layout["rows"]}
        )
    )
    updated = data(app.get_report_matrix(QUERY))
    assert updated["rows"][1]["archived"]
    assert value(updated, "2026-09-VARIANCE") == "-4000"
    assert not app.save_report_presentation(
        {
            "report_type": "SUBSIDIARY",
            "organization_id": "1",
            "expected_revision": matrix["matrix_revision"],
            "plans": {"2026-09": "2000"},
        }
    )["ok"]
    archived = updated["rows"][1]["cells"][1]
    assert not app.save_report_cells(
        {
            **QUERY,
            "base_revision": updated["matrix_revision"],
            "idempotency_key": uuid4().hex,
            "changes": [
                {
                    "coordinate": archived["coordinate"],
                    "value": {"kind": "QUANTITY", "quantity": "5"},
                }
            ],
        }
    )["ok"]


@pytest.mark.parametrize("plan", ["-1", "NaN", "Infinity", "1e100", "word"])
def test_invalid_plan(tmp_path: Path, plan: str) -> None:
    app = app_at(tmp_path)
    matrix = setup(app)
    assert not app.save_report_presentation(
        {
            "report_type": "SUBSIDIARY",
            "organization_id": "1",
            "expected_revision": matrix["matrix_revision"],
            "plans": {"2026-09": plan},
        }
    )["ok"]


def test_missing_supplier_receipt_explained_then_zero_recalculates(tmp_path: Path) -> None:
    app = app_at(tmp_path)
    matrix = setup(app)
    matrix = write(
        app,
        matrix,
        [
            (0, "2026-09-OPENING", "234"),
            (0, "2026-09-RECEIVED", "55"),
            (0, "2026-09-01", "3"),
            (0, "2026-09-07", "45"),
            (0, "2026-09-14", "65"),
        ],
    )
    cells = {c["column_id"]: c for c in matrix["rows"][0]["cells"]}
    assert "Производитель Б" in cells["2026-09-STOCK"]["issue"]["message"]
    assert cells["2026-09-STOCK"]["value"]["kind"] == "DATA_NOT_PROVIDED"
    matrix = write(app, matrix, [(1, "2026-09-RECEIVED", "0")])
    assert value(matrix, "2026-09-STOCK") == "176"
    assert value(matrix, "2026-09-VARIANCE") == "-3711"
    cells = {c["column_id"]: c for c in matrix["rows"][0]["cells"]}
    assert "issue" not in cells["2026-09-STOCK"]


@pytest.mark.parametrize("format_name", ["PNG", "JPEG"])
def test_subsidiary_pdf_embeds_detail_picture(tmp_path: Path, format_name: str) -> None:
    import base64
    import io

    from PIL import Image

    app = app_at(tmp_path)
    matrix = setup(app)
    picture = io.BytesIO()
    Image.new("RGB", (160, 80), "#21634e").save(picture, format=format_name)
    encoded = base64.b64encode(picture.getvalue()).decode("ascii")
    layout = data(app.get_report_layout({"report_type": "SUBSIDIARY", "organization_id": "1"}))
    layout["rows"][0]["configuration"]["image"] = (
        f"data:image/{'png' if format_name == 'PNG' else 'jpeg'};base64,{encoded}"
    )
    data(
        app.save_report_layout(
            {"report_type": "SUBSIDIARY", "organization_id": "1", "rows": layout["rows"]}
        )
    )
    matrix = data(app.get_report_matrix(QUERY))
    snapshot = monthly_snapshot(matrix, 9, "Организация")
    pdf = render_monthly_pdf(snapshot, {}, ROOT / "resources/fonts/ReportingSerif.ttf")
    assert b"/Subtype /Image" in pdf
    assert b"/Width 160" in pdf and b"/Height 80" in pdf


def test_pdf_shared_detail_continues_across_pages(tmp_path: Path) -> None:
    import copy
    import re

    app = app_at(tmp_path)
    matrix = setup(app)
    snapshot = monthly_snapshot(matrix, 9, "Организация")
    snapshot["rows"] = [snapshot["rows"][0]] + [
        copy.deepcopy(snapshot["rows"][1]) for _ in range(49)
    ]
    pdf = render_monthly_pdf(snapshot, {}, ROOT / "resources/fonts/ReportingSerif.ttf")
    pages = re.search(rb"/Count (\d+)", pdf)
    assert pages is not None and int(pages.group(1)) > 1


def test_export_opens_visible_month_and_keeps_values_and_import_map(tmp_path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.utils import get_column_letter

    app = app_at(tmp_path)
    matrix = setup(app)
    matrix = write(
        app,
        matrix,
        [
            (0, "2026-09-OPENING", "800"),
            (0, "2026-09-RECEIVED", "500"),
            (1, "2026-09-RECEIVED", "300"),
            (0, "2026-09-01", "560"),
            (1, "2026-09-01", "100"),
        ],
    )
    path = tmp_path / "visible-month.xlsx"
    app.configure_excel_dialogs(open_file=lambda: path, save_file=lambda _: path)
    data(
        app.export_report(
            {**QUERY, "visible_months": ["2026-09"], "stock_weeks": {"2026-09": "2026-09-01"}}
        )
    )
    workbook = load_workbook(path)
    sheet = workbook["Отчёт"]
    first = 7 + next(
        i for i, c in enumerate(matrix["time_columns"]) if c["id"] == "2026-09-OPENING"
    )
    letter = get_column_letter(first)
    assert sheet.sheet_view.pane is not None
    assert sheet.sheet_view.pane.topLeftCell == f"{letter}7"
    assert sheet.column_dimensions["G"].hidden
    assert not sheet.column_dimensions[letter].hidden
    assert sheet.cell(7, first).value == 800
    assert sheet.cell(7, first + 2).value == 940
    assert sheet.cell(7, first + 3).value == -2400
    assert f"{letter}7:{letter}8" in sheet.merged_cells
    assert workbook["Месячные планы"]["B2"].value == 1000
    assert workbook["_Системная карта"].max_row > 100
