"""Synthetic exchange safety: blanks, collisions, displacement, exact periods, formulas."""

from copy import deepcopy
from pathlib import Path
from typing import Any, cast

import pytest
from openpyxl import load_workbook

from backend.application.excel_reports import ParsedWorkbook
from backend.infrastructure.excel.matrix_exchange_v2 import read_exchange_snapshot
from backend.infrastructure.excel.openpyxl_matrix_workbook import OpenpyxlMatrixWorkbookAdapter
from backend.tests.integration.test_subsidiary_consumption import QUERY, app_at, data


def fixture(tmp_path: Path) -> tuple[Any, dict[str, Any], Path, OpenpyxlMatrixWorkbookAdapter]:
    app = app_at(tmp_path)
    matrix = data(app.get_report_matrix(QUERY))
    matrix["exchange_identity"] = {"dataset_id": "synthetic-dataset-a", "organization_id": "org-a"}
    path = tmp_path / "exchange.xlsx"
    adapter = OpenpyxlMatrixWorkbookAdapter()
    adapter.write(path, matrix)
    return app, matrix, path, adapter


def parse(
    adapter: OpenpyxlMatrixWorkbookAdapter, path: Path, matrix: dict[str, Any]
) -> ParsedWorkbook:
    return adapter.parse(path, report_type="SUBSIDIARY", organization_id=1, matrix=matrix)


def test_v2_preserves_exact_periods_typed_blank_and_export_metadata(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    snapshot = read_exchange_snapshot(path)
    assert snapshot and snapshot["exchange_identity"]["dataset_id"] == "synthetic-dataset-a"
    periods = [
        p for p in snapshot["periods"] if p["group_label"] == "2026-09" and p["kind"] == "USED"
    ]
    assert [(p["start"], p["end"]) for p in periods] == [
        ("2026-09-01", "2026-09-06"),
        ("2026-09-07", "2026-09-13"),
        ("2026-09-14", "2026-09-20"),
        ("2026-09-21", "2026-09-27"),
        ("2026-09-28", "2026-09-30"),
    ]
    parsed = parse(adapter, path, matrix)
    assert not parsed.issues
    assert parsed.skipped_count == len(parsed.cells)
    assert all(c.provenance["action"] == "KEEP" for c in parsed.cells)
    workbook = load_workbook(path)
    assert workbook["_Системная карта"]["B1"].value == 2
    assert workbook["_Системная карта"]["D6"].value == "DATA_NOT_PROVIDED"
    assert workbook["Месячные планы"]["A10"].value == "2026-09"


def test_explicit_clear_zero_decimal_comma_and_formula_have_distinct_actions(
    tmp_path: Path,
) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook["Отчёт"]
    sheet["BW7"] = "#CLEAR"
    sheet["BX7"] = 0
    sheet["CA7"] = "=80+50"
    sheet["CB7"] = "1 234,50"
    workbook.save(path)
    parsed = parse(adapter, path, matrix)
    assert not parsed.issues
    values = {cell.source_cell: cell for cell in parsed.cells}
    assert values["BW7"].provenance["action"] == "CLEAR"
    assert values["BW7"].value.kind == "DATA_NOT_PROVIDED"
    assert values["BX7"].value.quantity == "0"
    assert values["CA7"].value.quantity == "130"
    assert values["CA7"].provenance["formula"] == "=80+50"
    assert values["CB7"].value.quantity == "1234.5"


@pytest.mark.parametrize("formula", ["=1/0", "=UNSUPPORTED(1)", "='Other'!A1"])
def test_formula_error_never_becomes_zero(tmp_path: Path, formula: str) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    workbook["Отчёт"]["BW7"] = formula
    workbook.save(path)
    parsed = parse(adapter, path, matrix)
    assert any(i.code == "FORMULA_UNSUPPORTED" and i.source_cell == "BW7" for i in parsed.issues)
    assert all(c.source_cell != "BW7" for c in parsed.cells)


def test_foreign_database_with_same_local_ids_is_blocked(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    other = deepcopy(matrix)
    other["exchange_identity"]["dataset_id"] = "synthetic-dataset-b"
    assert any(i.code == "FOREIGN_DATASET" for i in parse(adapter, path, other).issues)


@pytest.mark.parametrize("change", ["row", "column"])
def test_inserted_rows_and_columns_cannot_reuse_stale_coordinate_map(
    tmp_path: Path, change: str
) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    if change == "row":
        workbook["Отчёт"].insert_rows(7)
    else:
        workbook["Отчёт"].insert_cols(7)
    workbook.save(path)
    assert any(i.code == "LAYOUT_CHANGED" for i in parse(adapter, path, matrix).issues)


def test_header_and_monthly_plans_are_staged_as_separate_actions(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    workbook["Шапка и выпуск"]["B1"] = "000.01"
    workbook["Месячные планы"]["B10"] = 120
    workbook.save(path)
    parsed = parse(adapter, path, matrix)
    assert not parsed.issues
    changes = cast(list[dict[str, Any]], parsed.metadata["structural_changes"])
    assert any(
        c["path"] == ["presentation", "header", "product_designation"] and c["after"] == "000.01"
        for c in changes
    )
    assert any(
        c["path"] == ["presentation", "plans", "2026-09"] and c["after"] == 120 for c in changes
    )


def test_unknown_hidden_sheet_requires_explicit_review(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook.create_sheet("Дополнительные факты")
    sheet.sheet_state = "hidden"
    sheet["A1"] = 123
    workbook.save(path)
    assert any(i.code == "SHEET_REVIEW_REQUIRED" for i in parse(adapter, path, matrix).issues)


def test_legacy_version_one_remains_readable(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    workbook["_Системная карта"]["B1"] = 1
    del workbook["_Обмен v2"]
    workbook["Отчёт"]["BW7"] = 100
    workbook["Отчёт"]["BX7"] = 20
    workbook["Отчёт"]["CA7"] = 5
    workbook.save(path)
    parsed = parse(adapter, path, matrix)
    assert not parsed.issues
    assert [c.value.quantity for c in parsed.cells if c.source_cell in {"BW7", "BX7", "CA7"}] == [
        "100",
        "20",
        "5",
    ]
    assert parsed.metadata["schema_version"] == 1


def test_legacy_id_collision_requires_matching_visible_business_context(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    workbook["_Системная карта"]["B1"] = 1
    del workbook["_Обмен v2"]
    workbook["Отчёт"]["C7"] = "Другая позиция с тем же локальным ID"
    workbook["Отчёт"]["BW7"] = 500
    workbook.save(path)
    parsed = parse(adapter, path, matrix)
    assert any(i.code == "LEGACY_CONTEXT_MISMATCH" for i in parsed.issues)


def test_weekly_supply_exchange_keeps_periods_and_calculated_monthly_receipts(
    tmp_path: Path,
) -> None:
    app = app_at(tmp_path)
    layout = data(app.get_report_layout({"report_type": "SUBSIDIARY", "organization_id": "1"}))
    layout["rows"][0]["configuration"]["subsidiary"]["weekly_supply"] = True
    data(
        app.save_report_layout(
            {"report_type": "SUBSIDIARY", "organization_id": "1", "rows": layout["rows"]}
        )
    )
    matrix = data(app.get_report_matrix(QUERY))
    path = tmp_path / "weekly_supply.xlsx"
    adapter = OpenpyxlMatrixWorkbookAdapter()
    adapter.write(path, matrix)
    snapshot = read_exchange_snapshot(path)
    assert snapshot
    weekly = [
        p for p in snapshot["periods"] if p["group_label"] == "2026-09" and p["kind"] == "SUPPLIED"
    ]
    assert [(p["start"], p["end"]) for p in weekly] == [
        ("2026-09-01", "2026-09-06"),
        ("2026-09-07", "2026-09-13"),
        ("2026-09-14", "2026-09-20"),
        ("2026-09-21", "2026-09-27"),
        ("2026-09-28", "2026-09-30"),
    ]
    metric_codes = {item["coordinate"]["metric_code"] for item in snapshot["coordinates"]}
    assert "SUB_SUPPLIED_PRIMARY" in metric_codes
    assert "SUB_RECEIVED_PRIMARY" not in metric_codes
    assert not parse(adapter, path, matrix).issues


def test_v2_import_subset_does_not_require_unrelated_target_rows(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    target = deepcopy(matrix)
    unrelated = deepcopy(target["rows"][0])
    unrelated["id"] = "unrelated-position"
    for cell in unrelated["cells"]:
        coordinate = cell["coordinate"]
        subject = "component_id" if coordinate.get("component_id") else "product_id"
        coordinate[subject] = "99999"
    target["rows"].append(unrelated)
    parsed = parse(adapter, path, target)
    assert not parsed.issues
    assert len(parsed.cells) == 79


def test_v2_deleted_map_entry_is_detected_from_original_manifest(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    workbook["_Системная карта"].delete_rows(6)
    workbook.save(path)
    parsed = parse(adapter, path, matrix)
    assert sum(i.code == "MISSING_COORDINATE" for i in parsed.issues) == 1


def test_v2_extra_or_changed_map_entry_is_not_accepted(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    mapping = workbook["_Системная карта"]
    mapping["A6"] = "BY7"  # calculated cell cannot be substituted for original input
    workbook.save(path)
    parsed = parse(adapter, path, matrix)
    assert any(i.code == "UNDECLARED_COORDINATE" for i in parsed.issues)
    assert any(i.code == "MISSING_COORDINATE" for i in parsed.issues)


@pytest.mark.parametrize(
    ("decision", "expected_error"),
    [
        ({"include": False, "reason": "Дублирующая аналитическая справка"}, None),
        ({"include": False, "reason": "  "}, "SHEET_EXCLUSION_REASON_REQUIRED"),
        ({"include": True, "reason": "Нужны значения"}, "UNSUPPORTED_SHEET"),
        ({"include": "false", "reason": "Справка"}, "SHEET_DECISION_INVALID"),
    ],
)
def test_unknown_sheet_decision_is_explicit_and_recorded(
    tmp_path: Path, decision: dict[str, Any], expected_error: str | None
) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    workbook = load_workbook(path)
    sheet = workbook.create_sheet("Дополнительный лист")
    sheet.sheet_state = "hidden"
    sheet["A1"] = 42
    workbook.save(path)
    initial = parse(adapter, path, matrix)
    description = next(
        s
        for s in cast(list[dict[str, Any]], initial.metadata["sheets"])
        if s["name"] == sheet.title
    )
    assert description == {
        "name": sheet.title,
        "state": "hidden",
        "known": False,
        "requires_decision": True,
        "included": None,
        "reason": "",
    }
    matrix["exchange_sheet_decisions"] = {sheet.title: decision}
    result = parse(adapter, path, matrix)
    if expected_error:
        assert [i.code for i in result.issues] == [expected_error]
    else:
        assert not result.issues
        saved = next(
            s
            for s in cast(list[dict[str, Any]], result.metadata["sheets"])
            if s["name"] == sheet.title
        )
        assert saved["included"] is False
        assert saved["reason"] == decision["reason"]


def test_sheet_exclusion_cannot_target_system_or_missing_sheet(tmp_path: Path) -> None:
    _, matrix, path, adapter = fixture(tmp_path)
    matrix["exchange_sheet_decisions"] = {
        "Отчёт": {"include": False, "reason": "Попытка исключения"},
        "Отсутствует": {"include": False, "reason": "Устаревшее решение"},
    }
    parsed = parse(adapter, path, matrix)
    assert {i.code for i in parsed.issues} == {"SYSTEM_SHEET_REQUIRED", "UNKNOWN_SHEET"}


def test_bridge_unknown_sheet_can_be_reviewed_without_reopening_file_dialog(tmp_path: Path) -> None:
    from backend.tests.integration.test_excel_bridge_dev33 import SUB, source_report

    app, source = source_report(tmp_path / "source")
    before = data(app.get_report_matrix(SUB))
    book = load_workbook(source)
    extra = book.create_sheet("Архивная справка")
    extra.sheet_state = "hidden"
    extra["A1"] = 123
    book.save(source)
    book.close()
    first = data(app.validate_import(SUB))
    assert first["error_count"] == 1
    assert first["issues"][0]["code"] == "SHEET_REVIEW_REQUIRED"
    assert any(
        s["name"] == "Архивная справка" and s["state"] == "hidden"
        for s in first["metadata"]["sheets"]
    )

    def forbidden_dialog() -> Path:
        raise AssertionError("Повторная проверка не должна открывать диалог выбора файла")

    app.configure_excel_dialogs(open_file=forbidden_dialog, save_file=lambda _: source)
    reviewed = data(
        app.validate_import(
            {
                **SUB,
                "batch_id": first["batch_id"],
                "sheet_decisions": {
                    "Архивная справка": {
                        "include": False,
                        "reason": "Архивный контрольный снимок; значения дублируют основной лист",
                    }
                },
            }
        )
    )
    assert reviewed["error_count"] == 0, reviewed
    assert reviewed["file_name"] == source.name
    decision = next(s for s in reviewed["metadata"]["sheets"] if s["name"] == "Архивная справка")
    assert decision["included"] is False
    assert decision["reason"]
    committed = data(app.commit_import({"batch_id": reviewed["batch_id"], "year": 2026}))
    assert committed["status"] == "COMMITTED"
    assert committed["imported_count"] == 0
    assert data(app.get_report_matrix(SUB))["rows"] == before["rows"]
