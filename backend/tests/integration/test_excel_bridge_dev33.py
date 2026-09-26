"""End-to-end exchange acceptance at the same bridge boundary used by the UI."""

from __future__ import annotations

import hashlib
import json
from contextlib import closing
from pathlib import Path
from typing import Any, cast
from uuid import uuid4

import pytest
from openpyxl import Workbook, load_workbook

from backend.api.working_reference_bridge import WorkingReferenceApplicationBridge
from backend.infrastructure.database.migrator import connect_sqlite

ROOT = Path(__file__).resolve().parents[3]
SUB = {"report_type": "SUBSIDIARY", "organization_id": "1", "year": 2026}


def data(response: dict[str, Any]) -> dict[str, Any]:
    assert response["ok"] is True, response
    return cast(dict[str, Any], response["data"])


def application(root: Path) -> WorkingReferenceApplicationBridge:
    root.mkdir(parents=True, exist_ok=True)
    return WorkingReferenceApplicationBridge(
        root / "data" / "reporting.sqlite3",
        migrations_directory=ROOT / "backend/migrations",
        definitions_directory=ROOT / "resources/report-definitions",
        backups_directory=root / "backups",
        inbox_directory=root / "imports" / "inbox",
    )


def dialogs(app: WorkingReferenceApplicationBridge, source: Path) -> None:
    app.configure_excel_dialogs(open_file=lambda: source, save_file=lambda _name: source)


def save_cell(
    app: WorkingReferenceApplicationBridge, query: dict[str, Any], cell: dict[str, Any], value: str
) -> None:
    matrix = data(app.get_report_matrix(query))
    data(
        app.save_report_cells(
            {
                **query,
                "base_revision": matrix["matrix_revision"],
                "idempotency_key": uuid4().hex,
                "changes": [
                    {
                        "coordinate": cell["coordinate"],
                        "value": {"kind": "QUANTITY", "quantity": value},
                    }
                ],
            }
        )
    )


def source_report(root: Path) -> tuple[WorkingReferenceApplicationBridge, Path]:
    app = application(root)
    layout_query = {"report_type": "SUBSIDIARY", "organization_id": "1"}
    layout = data(app.get_report_layout(layout_query))
    row = layout["rows"][0]
    row["position_name"] = "Деталь 000.010"
    row["party_name"] = "Изготовитель А"
    row["configuration"]["norm"] = "2"
    row["configuration"]["subsidiary"].update(
        designation="000.010",
        number="01.02",
        suppliers=[
            {"id": "PRIMARY", "name": "Изготовитель А", "contract": "250", "archived": False}
        ],
    )
    data(app.save_report_layout({**layout_query, "rows": layout["rows"]}))
    data(
        app.save_report_presentation(
            {
                **layout_query,
                "plans": {"2026-09": "100"},
                "expected_revision": data(app.get_report_matrix(SUB))["matrix_revision"],
                "header": {
                    "product_name": "Тестовое изделие",
                    "product_designation": "ИЗД.001",
                    "factory_name": "Завод",
                },
            }
        )
    )
    matrix = data(app.get_report_matrix(SUB))
    first = matrix["rows"][0]
    opening = next(
        c
        for c in first["cells"]
        if c["coordinate"].get("metric_code") == "SUB_OPENING"
        and c["coordinate"].get("period_start") == "2026-09-01"
    )
    save_cell(app, SUB, opening, "100")
    received = next(
        c
        for c in first["cells"]
        if c["coordinate"].get("metric_code") == "SUB_RECEIVED_PRIMARY"
        and c["coordinate"].get("period_start") == "2026-09-01"
    )
    save_cell(app, SUB, received, "20")
    file = root / "source.xlsx"
    dialogs(app, file)
    data(app.export_report(SUB))
    return app, file


@pytest.mark.parametrize("report", ["DAILY_MOVEMENT", "HEAD_SITE", "SUBSIDIARY"])
def test_canonical_v2_bridge_roundtrip_same_database_and_reopen(
    tmp_path: Path, report: str
) -> None:
    app = application(tmp_path / report)
    query = {**SUB, "report_type": report}
    file = tmp_path / f"{report}.xlsx"
    dialogs(app, file)
    data(app.export_report(query))
    book = load_workbook(file)
    assert "_Обмен v2" in book.sheetnames
    address = str(book["_Системная карта"]["A6"].value)
    book["Отчёт"][address] = 37.5
    book.save(file)
    book.close()
    preview = data(app.validate_import(query))
    assert preview["error_count"] == 0, preview
    result = data(app.commit_import({"batch_id": preview["batch_id"], "year": 2026}))
    assert result["status"] == "COMMITTED"
    reopened = application(tmp_path / report)
    matrix = data(reopened.get_report_matrix(query))
    assert any(c["value"].get("quantity") == "37.5" for r in matrix["rows"] for c in r["cells"])
    again = data(app.validate_import(query))
    assert again["already_imported"] is True
    assert again["new_count"] == again["changed_count"] == 0
    assert again["same_count"] == 1
    repeated_commit = data(app.commit_import({"batch_id": preview["batch_id"], "year": 2026}))
    assert repeated_commit["already_committed"] is True
    assert repeated_commit["imported_count"] == 0
    assert repeated_commit["same_count"] == 1
    with closing(connect_sqlite(app._database_path)) as connection:
        assert connection.execute("SELECT count(*) FROM report_fact_revisions").fetchone() == (1,)
        assert connection.execute("SELECT status FROM exchange_import_packages").fetchone() == (
            "COMMITTED",
        )


def test_canonical_foreign_database_imports_structure_header_plans_and_facts(
    tmp_path: Path,
) -> None:
    _source, file = source_report(tmp_path / "source")
    target = application(tmp_path / "target")
    before = data(target.get_report_matrix(SUB))
    default_id = before["rows"][0]["workspace_id"]
    dialogs(target, file)
    preview = data(target.validate_import({**SUB, "mode": "update"}))
    assert preview["error_count"] == 0, preview
    assert preview["position_count"] >= 1
    data(target.commit_import({"batch_id": preview["batch_id"], "year": 2026}))
    reopened = application(tmp_path / "target")
    matrix = data(reopened.get_report_matrix(SUB))
    imported = next(r for r in matrix["rows"] if r["left_values"].get("designation") == "000.010")
    assert imported["workspace_id"] != default_id
    untouched = next(r for r in matrix["rows"] if r["workspace_id"] == default_id)
    assert all(
        c["value"]["kind"] == "DATA_NOT_PROVIDED"
        for c in untouched["cells"]
        if c["state"]["access"] == "editable"
    )
    assert imported["left_values"]["position"] == "Деталь 000.010"
    assert imported["left_values"]["party"] == "Изготовитель А"
    values = {
        c["coordinate"]["metric_code"]: c["value"].get("quantity")
        for c in imported["cells"]
        if c["coordinate"].get("period_start") == "2026-09-01"
        and c["state"]["access"] == "editable"
    }
    assert values["SUB_OPENING"] == "100"
    assert (
        next(value for metric, value in values.items() if metric.startswith("SUB_RECEIVED_"))
        == "20"
    )
    assert matrix["presentation"]["plans"]["2026-09"] == "100"
    assert matrix["presentation"]["header"]["product_designation"] == "ИЗД.001"
    with closing(connect_sqlite(reopened._database_path)) as connection:
        row = connection.execute(
            "SELECT source_sha256, metadata_json FROM import_batches WHERE id=?",
            (preview["batch_id"],),
        ).fetchone()
        metadata = json.loads(row[1])
        assert metadata["original_source_sha256"] == hashlib.sha256(file.read_bytes()).hexdigest()
        assert metadata["import_context"]["year"] == 2026
        assert metadata["import_context"]["target_identity"].get("dataset_id")
        assert row[0] != metadata["original_source_sha256"]


def test_canonical_stale_projection_rolls_back_structure_and_facts(tmp_path: Path) -> None:
    _source, file = source_report(tmp_path / "source")
    target = application(tmp_path / "target")
    dialogs(target, file)
    preview = data(target.validate_import(SUB))
    assert preview["error_count"] == 0, preview
    data(target.rename_organization({"organization_id": "1", "name": "Изменено после проверки"}))
    result = target.commit_import({"batch_id": preview["batch_id"], "year": 2026})
    assert result["ok"] is False, result
    assert cast(dict[str, Any], result["error"])["code"] == "REVISION_CONFLICT"
    matrix = data(target.get_report_matrix(SUB))
    assert not any(r["left_values"].get("designation") == "000.010" for r in matrix["rows"])
    with closing(connect_sqlite(target._database_path)) as connection:
        assert connection.execute("SELECT count(*) FROM report_fact_revisions").fetchone() == (0,)
        assert connection.execute(
            "SELECT status FROM import_batches WHERE id=?", (preview["batch_id"],)
        ).fetchone() == ("STAGED",)
        assert connection.execute(
            "SELECT status FROM exchange_import_packages WHERE id=?", (preview["batch_id"],)
        ).fetchone() == ("STAGED",)


def enterprise_workbook(path: Path) -> None:
    book = Workbook()
    sheet = book.active
    assert sheet is not None
    sheet["A1"] = "Отчёт 2026"
    for address, value in {
        "B6": "№ п/п",
        "C6": "Обозначение",
        "D6": "Наименование",
        "E6": "Входимость",
        "F6": "Производитель",
        "H6": "Объём поставок по договору",
        "CU6": "Факт",
        "CU2": "Сентябрь 2026",
        "CU3": "1 неделя",
        "C4": "=1/2",
        "B8": "01.1",
        "C8": "000.010",
        "D8": "Деталь",
        "E8": 2,
        "F8": "Первый",
        "H8": 30,
        "CU8": "=80+50",
        "F9": "Второй",
        "H9": 20,
        "CU9": 0,
    }.items():
        sheet[address] = cast(str | int, value)
    sheet["C4"].number_format = "0.0%"
    for column in "BCDE":
        sheet.merge_cells(f"{column}8:{column}9")
    sheet.row_dimensions[9].hidden = True
    sheet.column_dimensions["CU"].hidden = True
    book.save(path)
    book.close()


def test_enterprise_bridge_auto_maps_suppliers_and_explicit_week_without_manual_cells(
    tmp_path: Path,
) -> None:
    app = application(tmp_path / "app")
    file = tmp_path / "enterprise.xlsx"
    enterprise_workbook(file)
    dialogs(app, file)
    raw = data(app.validate_import(SUB))
    identity = raw["reference_workbook"]["id"]
    request = {
        **SUB,
        "action": "transfer",
        "id": identity,
        "mode": "update",
        "mappings": [],
        "period_rules": {
            "0:CU": {
                "start": "2026-09-01",
                "end": "2026-09-06",
                "reason": "Проверен календарь предприятия",
            }
        },
    }
    unresolved = data(app.reference_report({**request, "period_rules": {}}))
    assert unresolved["error_count"] > 0
    assert any(issue["code"] == "PERIOD_REQUIRED" for issue in unresolved["issues"])
    with closing(connect_sqlite(app._database_path)) as connection:
        assert connection.execute("SELECT count(*) FROM report_fact_revisions").fetchone() == (0,)
    preview = data(app.reference_report(request))
    assert preview["error_count"] == 0, preview
    assert preview["new_count"] == 2
    assert len(preview["auto_mappings"]) == 2
    data(app.commit_import({"batch_id": preview["batch_id"], "year": 2026}))
    reopened = application(tmp_path / "app")
    matrix = data(reopened.get_report_matrix(SUB))
    rows = [r for r in matrix["rows"] if r["left_values"].get("designation") == "000.010"]
    assert len(rows) == 2
    assert {r["left_values"]["party"] for r in rows} == {"Первый", "Второй"}
    observed = {}
    for row in rows:
        cell = next(
            c
            for c in row["cells"]
            if c["coordinate"].get("metric_code", "").startswith("SUB_SUPPLIED_")
            and c["coordinate"].get("period_start") == "2026-09-01"
        )
        observed[row["left_values"]["party"]] = cell["value"].get("quantity")
    assert observed == {"Первый": "130", "Второй": "0"}
    repeated = data(app.reference_report(request))
    assert repeated["already_imported"] is True
    assert repeated["new_count"] == repeated["changed_count"] == 0
    assert repeated["same_count"] == 2
    with closing(connect_sqlite(app._database_path)) as connection:
        assert connection.execute("SELECT count(*) FROM report_fact_revisions").fetchone() == (2,)
        assert connection.execute(
            "SELECT count(*) FROM reference_transfer_decisions"
        ).fetchone() == (1,)
        metadata = json.loads(
            connection.execute(
                "SELECT metadata_json FROM import_batches WHERE id=?", (preview["batch_id"],)
            ).fetchone()[0]
        )
        assert metadata["import_context"]["year"] == 2026
        assert metadata["import_context"]["target_identity"].get("dataset_id")


def test_canonical_foreign_create_exact_repeat_is_successful_noop(tmp_path: Path) -> None:
    _source, file = source_report(tmp_path / "source")
    target = application(tmp_path / "target")
    dialogs(target, file)
    first = data(target.validate_import({**SUB, "mode": "create"}))
    assert first["error_count"] == 0, first
    assert first["new_count"] == 2
    data(target.commit_import({"batch_id": first["batch_id"], "year": 2026}))
    repeated = data(target.validate_import({**SUB, "mode": "create"}))
    assert repeated["already_imported"] is True, repeated
    assert repeated["error_count"] == 0
    assert repeated["new_count"] == repeated["changed_count"] == 0
    assert repeated["same_count"] == 2
    committed = data(target.commit_import({"batch_id": repeated["batch_id"], "year": 2026}))
    assert committed["already_committed"] is True
    assert committed["imported_count"] == 0
    with closing(connect_sqlite(target._database_path)) as connection:
        assert connection.execute("SELECT count(*) FROM report_fact_revisions").fetchone() == (2,)
        assert connection.execute("SELECT count(*) FROM import_batches").fetchone() == (1,)


def test_enterprise_create_repeat_is_noop_but_changed_source_cannot_overwrite(
    tmp_path: Path,
) -> None:
    app = application(tmp_path / "app")
    file = tmp_path / "enterprise.xlsx"
    enterprise_workbook(file)
    dialogs(app, file)
    raw = data(app.validate_import({**SUB, "mode": "create"}))
    request = {
        **SUB,
        "action": "transfer",
        "id": raw["reference_workbook"]["id"],
        "mode": "create",
        "mappings": [],
        "period_rules": {
            "0:CU": {"start": "2026-09-01", "end": "2026-09-06", "reason": "Проверен календарь"}
        },
    }
    first = data(app.reference_report(request))
    assert first["error_count"] == 0, first
    data(app.commit_import({"batch_id": first["batch_id"], "year": 2026}))
    repeated = data(app.reference_report(request))
    assert repeated["already_imported"] is True, repeated
    assert repeated["error_count"] == 0
    assert repeated["new_count"] == repeated["changed_count"] == 0
    assert repeated["same_count"] == 2
    again = data(app.commit_import({"batch_id": repeated["batch_id"], "year": 2026}))
    assert again["already_committed"] is True and again["imported_count"] == 0
    book = load_workbook(file)
    book.worksheets[0]["CU8"] = 999
    book.save(file)
    book.close()
    changed_raw = data(app.validate_import({**SUB, "mode": "create"}))
    changed = data(app.reference_report({**request, "id": changed_raw["reference_workbook"]["id"]}))
    assert changed["error_count"] > 0, changed
    assert changed["transfer_validated"] is False
    assert any(issue["code"] == "STRUCTURE_REQUIRED" for issue in changed["issues"])
    with closing(connect_sqlite(app._database_path)) as connection:
        assert connection.execute(
            "SELECT quantity FROM report_fact_revisions ORDER BY id"
        ).fetchall() == [("130",), ("0",)]
        assert connection.execute("SELECT count(*) FROM import_batches").fetchone() == (1,)


def test_missing_canonical_stage_package_cannot_commit_unreviewed_facts(tmp_path: Path) -> None:
    _source, file = source_report(tmp_path / "source")
    target = application(tmp_path / "target")
    dialogs(target, file)
    preview = data(target.validate_import(SUB))
    assert preview["error_count"] == 0, preview
    with closing(connect_sqlite(target._database_path)) as connection, connection:
        connection.execute(
            "DELETE FROM exchange_import_packages WHERE id=?", (preview["batch_id"],)
        )
    result = target.commit_import({"batch_id": preview["batch_id"], "year": 2026})
    assert result["ok"] is False
    assert "пакет структуры отсутствует" in cast(dict[str, Any], result["error"])["message"]
    with closing(connect_sqlite(target._database_path)) as connection:
        assert connection.execute("SELECT count(*) FROM report_fact_revisions").fetchone() == (0,)
        assert connection.execute(
            "SELECT status FROM import_batches WHERE id=?", (preview["batch_id"],)
        ).fetchone() == ("STAGED",)
    matrix = data(target.get_report_matrix(SUB))
    assert not any(row["left_values"].get("designation") == "000.010" for row in matrix["rows"])
