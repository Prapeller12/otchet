from __future__ import annotations

import base64
import io
import json
import shutil
from datetime import date
from pathlib import Path
from typing import Any
from uuid import uuid4

import pytest
from openpyxl import load_workbook
from PIL import Image

from backend.api.working_reference_bridge import WorkingReferenceApplicationBridge
from backend.infrastructure.database.migrator import apply_migrations, connect_sqlite

ROOT = Path(__file__).resolve().parents[3]
QUERY = {"report_type": "DAILY_MOVEMENT", "organization_id": "1"}


def data(response: dict[str, Any]) -> Any:
    assert response["ok"], response
    return response["data"]


@pytest.fixture
def app(tmp_path: Path) -> WorkingReferenceApplicationBridge:
    database = tmp_path / "reporting.db"
    connection = connect_sqlite(database)
    apply_migrations(connection, ROOT / "backend/migrations")
    connection.close()
    return WorkingReferenceApplicationBridge(
        database,
        migrations_directory=ROOT / "backend/migrations",
        definitions_directory=ROOT / "resources/report-definitions",
        backups_directory=tmp_path / "backups",
        inbox_directory=tmp_path / "inbox",
    )


def picture() -> str:
    stream = io.BytesIO()
    Image.new("RGB", (8, 8), "blue").save(stream, "PNG")
    return "data:image/png;base64," + base64.b64encode(stream.getvalue()).decode()


def test_default_balance_midmonth_preview_save_and_month_carry(
    app: WorkingReferenceApplicationBridge, tmp_path: Path
) -> None:
    query = {**QUERY, "year": 2026}
    matrix = data(app.get_report_matrix(query))
    assert len(matrix["time_columns"]) == 365
    indices = {column["id"]: i for i, column in enumerate(matrix["time_columns"])}
    rows = {row["metric_code"]: row for row in matrix["rows"]}
    changes = [
        {
            "coordinate": rows[code]["cells"][indices[day]]["coordinate"],
            "value": {"kind": "QUANTITY", "quantity": quantity},
        }
        for code, day, quantity in [
            ("WRK_DAILY_RECEIVED", "2026-09-14", "20"),
            ("WRK_DAILY_USED", "2026-09-14", "3"),
            ("WRK_DAILY_USED", "2026-10-01", "2"),
        ]
    ]
    preview = data(app.get_report_matrix({**query, "preview_changes": changes}))
    balance = next(row for row in preview["rows"] if row["metric_code"] == "WRK_DAILY_BALANCE")
    assert balance["cells"][indices["2026-09-13"]]["value"] == {"kind": "DATA_NOT_PROVIDED"}
    assert balance["cells"][indices["2026-09-14"]]["value"] == {
        "kind": "QUANTITY",
        "quantity": "17",
    }
    assert balance["cells"][indices["2026-10-01"]]["value"] == {
        "kind": "QUANTITY",
        "quantity": "15",
    }
    fresh = data(app.get_report_matrix(query))
    assert fresh == matrix  # Preview is read-only, including the fact revision.
    data(
        app.save_report_cells(
            {
                **query,
                "base_revision": matrix["matrix_revision"],
                "idempotency_key": uuid4().hex,
                "changes": changes,
            }
        )
    )
    saved = data(app.get_report_matrix(query))
    assert saved["rows"][2]["cells"] == balance["cells"]
    assert saved["rows"][0]["cells"][0]["value"] == {"kind": "DATA_NOT_PROVIDED"}
    conn = connect_sqlite(tmp_path / "reporting.db")
    try:
        assert conn.execute("SELECT count(*) FROM report_fact_revisions").fetchone()[0] == 3
    finally:
        conn.close()


def test_year_scope_leap_day_and_invalid_preview(app: WorkingReferenceApplicationBridge) -> None:
    matrix = data(app.get_report_matrix({**QUERY, "year": 2024}))
    assert len(matrix["time_columns"]) == 366
    assert any(column["id"] == "2024-02-29" for column in matrix["time_columns"])
    for year in [True, "2026", 1899, 2101]:
        assert not app.get_report_matrix({**QUERY, "year": year})["ok"]
    protected = matrix["rows"][2]["cells"][0]["coordinate"]
    assert not app.get_report_matrix(
        {
            **QUERY,
            "year": 2024,
            "preview_changes": [
                {"coordinate": protected, "value": {"kind": "QUANTITY", "quantity": "10"}}
            ],
        }
    )["ok"]


def test_report_presentation_persists_and_is_scoped(app: WorkingReferenceApplicationBridge) -> None:
    data(
        app.save_report_presentation(
            {**QUERY, "title": "Мой ежедневный отчёт", "widths": {"party": 123, "2026-09-14": 99}}
        )
    )
    data(app.save_report_presentation({**QUERY, "widths": {"party": 210}}))
    matrix = data(app.get_report_matrix({**QUERY, "year": 2026}))
    assert matrix["title"] == "Мой ежедневный отчёт"
    assert matrix["presentation"]["widths"] == {"party": 210}
    other = data(app.get_report_matrix({"report_type": "HEAD_SITE", "organization_id": "1"}))
    assert other["title"] != matrix["title"]
    patches: list[dict[str, Any]] = [
        {"title": " "},
        {"title": "x" * 201},
        {"widths": {"party": 1}},
        {"widths": {"party": True}},
    ]
    for patch in patches:
        assert not app.save_report_presentation({**QUERY, **patch})["ok"]


def test_upgrade_preserves_opening_date_and_repairs_previous_preset(tmp_path: Path) -> None:
    old_migrations = tmp_path / "old-migrations"
    old_migrations.mkdir()
    for path in (ROOT / "backend/migrations").glob("*.sql"):
        if path.name < "0007":
            shutil.copyfile(path, old_migrations / path.name)
    database = tmp_path / "upgrade.db"
    conn = connect_sqlite(database)
    apply_migrations(conn, old_migrations)
    conn.close()
    app = WorkingReferenceApplicationBridge(
        database,
        migrations_directory=old_migrations,
        definitions_directory=ROOT / "resources/report-definitions",
        backups_directory=tmp_path / "backups",
    )
    layout = data(app.get_report_layout(QUERY))
    config = layout["rows"][0]["configuration"]
    config.update(opening="5", norm="3")
    config["indicators"][2]["formula"] = "=OPENING+CUM(WRK_DAILY_RECEIVED)-CUM(WRK_DAILY_USED)"
    config["indicators"].append(layout["presets"][0]["indicators"][1])
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    conn = connect_sqlite(database)
    try:
        assert apply_migrations(conn, ROOT / "backend/migrations") == (
            "0007",
            "0008",
            "0009",
            "0010",
            "0011",
        )
        assert apply_migrations(conn, ROOT / "backend/migrations") == ()
    finally:
        conn.close()
    updated = data(app.get_report_layout(QUERY))["rows"][0]["configuration"]
    assert updated["opening"] == "5"
    assert updated["opening_date"] == date.today().replace(day=1).isoformat()
    assert updated["indicators"][2]["formula"] == "=BALANCE(WRK_DAILY_RECEIVED,WRK_DAILY_USED)"
    year = data(app.get_report_matrix({**QUERY, "year": date.today().year}))
    index = next(
        i for i, c in enumerate(year["time_columns"]) if c["id"] == updated["opening_date"]
    )
    assert year["rows"][2]["cells"][index]["value"] == {"kind": "QUANTITY", "quantity": "5"}
    if index:
        assert year["rows"][2]["cells"][index - 1]["value"] == {"kind": "DATA_NOT_PROVIDED"}


@pytest.mark.parametrize("report", ["DAILY_MOVEMENT", "HEAD_SITE", "SUBSIDIARY"])
def test_year_excel_roundtrip(
    app: WorkingReferenceApplicationBridge, tmp_path: Path, report: str
) -> None:
    query = {"report_type": report, "organization_id": "1", "year": 2026}
    destination = tmp_path / "year.xlsx"
    matrix = data(app.get_report_matrix(query))
    first_row = next(
        row for row in matrix["rows"] if row["cells"][0]["state"]["access"] == "editable"
    )
    changes = [
        {
            "coordinate": [c for c in first_row["cells"] if c["state"]["access"] == "editable"][
                index
            ]["coordinate"],
            "value": {"kind": "QUANTITY", "quantity": value},
        }
        for index, value in [(0, "12.25"), (-1, "0")]
    ]
    data(
        app.save_report_cells(
            {
                **query,
                "base_revision": matrix["matrix_revision"],
                "idempotency_key": uuid4().hex,
                "changes": changes,
            }
        )
    )
    app.configure_excel_dialogs(open_file=lambda: destination, save_file=lambda _: destination)
    data(app.export_report(query))
    preview = data(app.validate_import(query))
    assert preview["error_count"] == 0
    data(app.commit_import({"batch_id": preview["batch_id"], "year": 2026}))
    saved = data(app.get_report_matrix(query))
    for index, value in [(0, "12.25"), (-1, "0")]:
        assert [c for c in saved["rows"][0]["cells"] if c["state"]["access"] == "editable"][index][
            "value"
        ] == {"kind": "QUANTITY", "quantity": value}


@pytest.mark.parametrize(
    "report,preset_label,input_code,output_code,expected",
    [
        (
            "DAILY_MOVEMENT",
            "Наличие и комплектность — как в Excel",
            "AVAILABLE_QTY",
            "READY_SETS",
            "6",
        ),
    ],
)
def test_source_presets_without_manual_formulas(
    app: WorkingReferenceApplicationBridge,
    report: str,
    preset_label: str,
    input_code: str,
    output_code: str,
    expected: str,
) -> None:
    query = {"report_type": report, "organization_id": "1"}
    layout = data(app.get_report_layout(query))
    preset = next(item for item in layout["presets"] if item["label"] == preset_label)
    position = next(
        row
        for row in layout["rows"]
        if set(preset["required_codes"])
        <= {field["code"] for field in row["configuration"]["indicators"]}
    )
    config = position["configuration"]
    config["norm"] = "3"
    indicators = {item["code"]: item for item in config["indicators"]}
    indicators.update({item["code"]: item for item in preset["indicators"]})
    config["indicators"] = list(indicators.values())
    data(app.save_report_layout({**query, "rows": layout["rows"]}))
    matrix = data(app.get_report_matrix({**query, "year": 2026}))
    rows = {
        row["metric_code"]: row
        for row in matrix["rows"]
        if row["group_id"] == f"workspace-group-{position['id']}"
    }
    changes = [
        {
            "coordinate": rows[input_code]["cells"][13]["coordinate"],
            "value": {"kind": "QUANTITY", "quantity": "20"},
        }
    ]
    if "CONTRACT_QTY" in rows:
        changes.append(
            {
                "coordinate": rows["CONTRACT_QTY"]["cells"][13]["coordinate"],
                "value": {"kind": "QUANTITY", "quantity": "100"},
            }
        )
    preview = data(app.get_report_matrix({**query, "year": 2026, "preview_changes": changes}))
    actual = next(
        row
        for row in preview["rows"]
        if row["group_id"] == f"workspace-group-{position['id']}"
        and row["metric_code"] == output_code
    )
    assert actual["cells"][13]["value"] == {"kind": "QUANTITY", "quantity": expected}


def save_values(
    app: WorkingReferenceApplicationBridge, values: dict[str, list[str | None]]
) -> None:
    matrix = data(app.get_report_matrix(QUERY))
    changes = []
    for row in matrix["rows"]:
        for index, value in enumerate(values.get(row["metric_code"], [])):
            changes.append(
                {
                    "coordinate": row["cells"][index]["coordinate"],
                    "value": {"kind": "DATA_NOT_PROVIDED"}
                    if value is None
                    else {"kind": "QUANTITY", "quantity": value},
                }
            )
    data(
        app.save_report_cells(
            {
                **QUERY,
                "base_revision": matrix["matrix_revision"],
                "idempotency_key": uuid4().hex,
                "changes": changes,
            }
        )
    )


def test_configuration_images_formulas_and_history_survive_restart(
    app: WorkingReferenceApplicationBridge, tmp_path: Path
) -> None:
    layout = data(app.get_report_layout(QUERY))
    save_values(app, {"WRK_DAILY_RECEIVED": ["71", "19"], "WRK_DAILY_USED": ["0", "0"]})
    position = layout["rows"][0]
    config = position["configuration"]
    config.update(category="DSE", image=picture(), norm="3", opening="0")
    config["indicators"][2]["formula"] = "=OPENING+CUM(WRK_DAILY_RECEIVED)-CUM(WRK_DAILY_USED)"
    config["indicators"].append(
        {
            "code": "READY_SETS",
            "label": "Комплектность",
            "formula": "=IF(WRK_DAILY_BALANCE>=NORM,ROUNDDOWN(WRK_DAILY_BALANCE/NORM,0),0)",
        }
    )
    position["position_name"] = "ДСЕ с фотографией"
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    assert list((tmp_path / "backups").glob("*.sqlite3"))
    assert data(app.get_report_layout(QUERY))["rows"][0] == position
    matrix = data(app.get_report_matrix(QUERY))
    assert matrix["rows"][0]["image"] == picture()
    assert matrix["rows"][0]["category"] == "DSE"
    ready = next(row for row in matrix["rows"] if row["metric_code"] == "WRK_DAILY_READY_SETS")
    assert [cell["value"] for cell in ready["cells"][:2]] == [
        {"kind": "QUANTITY", "quantity": "23"},
        {"kind": "QUANTITY", "quantity": "30"},
    ]
    assert ready["cells"][2]["value"] == {"kind": "DATA_NOT_PROVIDED"}
    denied = app.save_report_cells(
        {
            **QUERY,
            "base_revision": matrix["matrix_revision"],
            "idempotency_key": uuid4().hex,
            "changes": [
                {
                    "coordinate": ready["cells"][0]["coordinate"],
                    "value": {"kind": "QUANTITY", "quantity": "999"},
                }
            ],
        }
    )
    assert not denied["ok"]
    # Removing/re-adding a configured indicator resurfaces the same immutable fact.
    removed = config["indicators"].pop(0)
    config["indicators"] = [
        item for item in config["indicators"] if item["code"] == "WRK_DAILY_USED"
    ]
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    config["indicators"].insert(0, removed)
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    restarted = WorkingReferenceApplicationBridge(
        tmp_path / "reporting.db",
        migrations_directory=ROOT / "backend/migrations",
        definitions_directory=ROOT / "resources/report-definitions",
    )
    assert (
        data(restarted.get_report_matrix(QUERY))["rows"][0]["cells"][0]["value"]["quantity"] == "71"
    )


@pytest.mark.parametrize(
    "invalid",
    [
        {"category": []},
        {"category": "FAKE"},
        {"norm": "0"},
        {"norm": "NaN"},
        {"norm": "1e-999999999"},
        {"image": "data:image/png;base64,SGVsbG8="},
        {"image": "https://outside/image.png"},
        {"indicators": [{"code": "A", "label": "A", "formula": "=A+1"}]},
        {"indicators": [{"code": "A", "label": "A", "formula": "=UNKNOWN+1"}]},
        {"indicators": [{"code": "A", "label": "A", "formula": "=__import__('os')"}]},
        {"indicators": [{"code": "A", "label": "A", "formula": ""}] * 2},
    ],
)
def test_invalid_configuration_is_atomic(
    app: WorkingReferenceApplicationBridge, invalid: dict[str, Any]
) -> None:
    before = data(app.get_report_layout(QUERY))
    rows = json.loads(json.dumps(before["rows"]))
    rows[0]["position_name"] = "Must not persist"
    rows[1]["configuration"].update(invalid)
    assert not app.save_report_layout({**QUERY, "rows": rows})["ok"]
    assert data(app.get_report_layout(QUERY)) == before


def test_template_category_changes_and_foreign_rows(app: WorkingReferenceApplicationBridge) -> None:
    layout = data(app.get_report_layout(QUERY))
    position = layout["rows"][0]
    save_values(app, {"WRK_DAILY_RECEIVED": ["12"]})
    original_template = position["template_group_id"]
    position["template_group_id"] = layout["rows"][1]["template_group_id"]
    position["configuration"]["category"] = "PART"
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    position["template_group_id"] = original_template
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    assert data(app.get_report_matrix(QUERY))["rows"][0]["cells"][0]["value"]["quantity"] == "12"
    org = data(app.create_organization({"name": "Дочернее общество"}))["organization"]["id"]
    assert not app.save_report_layout({**QUERY, "organization_id": org, "rows": layout["rows"]})[
        "ok"
    ]


def test_custom_fields_excel_roundtrip_and_stale_import(
    app: WorkingReferenceApplicationBridge, tmp_path: Path
) -> None:
    layout = data(app.get_report_layout(QUERY))
    config = layout["rows"][0]["configuration"]
    config["indicators"].extend(
        [
            {"code": "EXTRA", "label": "Пользовательский ввод", "formula": ""},
            {"code": "DOUBLE", "label": "Удвоенное", "formula": "=EXTRA*2"},
        ]
    )
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    save_values(app, {"EXTRA": ["3.125"]})
    destination = tmp_path / "export.xlsx"
    app.configure_excel_dialogs(open_file=lambda: destination, save_file=lambda _: destination)
    data(app.export_report(QUERY))
    book = load_workbook(destination)
    sheet = book["Отчёт"]
    double_row = next(row for row in sheet.iter_rows() if row[2].value == "Удвоенное")
    assert double_row[4].value == 6.25
    assert double_row[4].protection.locked
    source_cell = next(
        row[0].value
        for row in book["_Системная карта"].iter_rows(min_row=6)
        if '"metric_code":"EXTRA"' in str(row[1].value)
    )
    sheet[str(source_cell)] = 8
    book.save(destination)
    preview = data(app.validate_import(QUERY))
    data(app.commit_import({"batch_id": preview["batch_id"]}))
    matrix = data(app.get_report_matrix(QUERY))
    double = next(row for row in matrix["rows"] if row["metric_code"] == "DOUBLE")
    assert double["cells"][0]["value"]["quantity"] == "16"
    # A validated input must not become writable after it is converted to a formula.
    sheet[str(source_cell)] = 9
    book.save(destination)
    stale = data(app.validate_import(QUERY))
    next(item for item in config["indicators"] if item["code"] == "EXTRA")["formula"] = "=5"
    data(app.save_report_layout({**QUERY, "rows": layout["rows"]}))
    assert not app.commit_import({"batch_id": stale["batch_id"]})["ok"]


def test_all_twelve_source_positions_produce_23_and_30(
    app: WorkingReferenceApplicationBridge,
) -> None:
    # Source: ежедневный отчёт.xlsx, Учёт ДСЕ и ПКИ, B1:M1 and B4:M5.
    norms = [1, 1, 3, 1, 1, 1, 2, 3, 4, 1, 1, 1]
    first = [200, 144, 71, 50, 130, 130, 133, 180, 500, 600, 200, 166]
    second = [1000, 500, 90, 87, 221, 569, 996, 108, 500, 800, 1200, 89]
    layout = data(app.get_report_layout(QUERY))
    component = layout["rows"][0]
    rows = []
    for index, norm in enumerate(norms):
        position = json.loads(json.dumps(component))
        position["id"] = None
        position["position_name"] = f"Контрольная позиция {index + 1}"
        config = position["configuration"]
        config.update(category="PKI", norm=str(norm), opening=str(first[index]))
        config["indicators"][2].update(layout["presets"][0]["indicators"][0])
        config["indicators"].append(layout["presets"][0]["indicators"][1])
        rows.append(position)
    data(app.save_report_layout({**QUERY, "rows": rows}))
    matrix = data(app.get_report_matrix(QUERY))
    changes = []
    for index in range(12):
        received, used = matrix["rows"][index * 4 : index * 4 + 2]
        # Reproduce the two supplied snapshots with separate received/used declarations.
        delta = second[index] - first[index]
        for row, amounts in ((received, [0, max(delta, 0)]), (used, [0, max(-delta, 0)])):
            for day, amount in enumerate(amounts):
                changes.append(
                    {
                        "coordinate": row["cells"][day]["coordinate"],
                        "value": {"kind": "QUANTITY", "quantity": str(amount)},
                    }
                )
    data(
        app.save_report_cells(
            {
                **QUERY,
                "base_revision": matrix["matrix_revision"],
                "idempotency_key": uuid4().hex,
                "changes": changes,
            }
        )
    )
    matrix = data(app.get_report_matrix(QUERY))
    ready = next(row for row in matrix["rows"] if row["metric_code"] == "WRK_DAILY_READY_SETS")
    assert [cell["value"]["quantity"] for cell in ready["cells"][:2]] == ["23", "30"]
