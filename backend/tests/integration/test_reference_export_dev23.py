"""Reference ordering and value fidelity without private source fixtures."""

from pathlib import Path

import pytest
from openpyxl import load_workbook
from reportlab import rl_config  # type: ignore[import-untyped]

from backend.application.monthly_report import monthly_snapshot
from backend.infrastructure.excel.openpyxl_matrix_workbook import OpenpyxlMatrixWorkbookAdapter
from backend.infrastructure.subsidiary_pdf import render_subsidiary_pdf
from backend.tests.integration.test_subsidiary_consumption import ROOT, app_at, data


def test_daily_calculated_total_is_closing_numeric_snapshot(tmp_path: Path) -> None:
    app = app_at(tmp_path)
    matrix = data(
        app.get_report_matrix(
            {"report_type": "DAILY_MOVEMENT", "organization_id": "1", "year": 2026}
        )
    )
    calculated = next(
        row
        for row in matrix["rows"]
        if row.get("indicator_detail", {}).get("kind") == "CALCULATION"
    )
    calculated["cells"][-1]["value"] = {"kind": "QUANTITY", "quantity": "12345"}
    position = matrix["rows"].index(calculated) + 7
    destination = tmp_path / "daily.xlsx"
    OpenpyxlMatrixWorkbookAdapter().write(destination, matrix)
    workbook = load_workbook(destination)
    cell = workbook["Отчёт"].cell(position, len(matrix["left_columns"]) + 1)
    assert cell.value == 12345
    assert cell.protection.locked
    assert cell.comment and "не суммируются" in cell.comment.text


def test_excel_headers_group_months_and_keep_literal_user_names(tmp_path: Path) -> None:
    app = app_at(tmp_path)
    matrix = data(
        app.get_report_matrix({"report_type": "HEAD_SITE", "organization_id": "1", "year": 2026})
    )
    matrix["rows"][0]["left_values"]["position"] = "=1+1"
    matrix["presentation"]["header"] = {
        "product_designation": "=2+2",
        "factory_name": "Тестовый завод",
    }
    matrix["presentation"]["production_codes"] = [
        {"id": "A", "label": "=3+3", "plans": {"2026-01": "12345"}, "actuals": {"2026-01": "0"}}
    ]
    matrix["presentation"]["production_code_annual"] = {"A": "54321"}
    destination = tmp_path / "head.xlsx"
    adapter = OpenpyxlMatrixWorkbookAdapter()
    count = adapter.write(destination, matrix)
    workbook = load_workbook(destination)
    assert workbook["Отчёт"]["C7"].data_type == "s"
    assert workbook["Шапка и выпуск"]["B1"].value == "=2+2"
    assert workbook["Шапка и выпуск"]["B1"].data_type == "s"
    code_row = next(row for row in workbook["Шапка и выпуск"] if row[0].value == "=3+3")
    assert [cell.value for cell in code_row[:4]] == ["=3+3", 54321, 12345, 0]
    assert any(
        r.min_row == 5 and r.max_col > r.min_col for r in workbook["Отчёт"].merged_cells.ranges
    )
    parsed = adapter.parse(destination, report_type="HEAD_SITE", organization_id=1, matrix=matrix)
    assert not parsed.issues
    assert len(parsed.cells) == count


def test_reference_pdf_includes_report_metadata_and_code_snapshot(
    tmp_path: Path, monkeypatch: pytest.MonkeyPatch
) -> None:
    monkeypatch.setattr(rl_config, "pageCompression", 0)
    app = app_at(tmp_path)
    matrix = data(
        app.get_report_matrix({"report_type": "HEAD_SITE", "organization_id": "1", "year": 2026})
    )
    snapshot = monthly_snapshot(matrix, 9, "Тестовая площадка")
    snapshot["header"] = {
        "product_designation": "SYNTHETIC-12345",
        "factory_name": "Тестовый завод",
    }
    snapshot["production_codes"] = [
        {
            "id": "A",
            "label": "TEST-CODE-A",
            "plans": {"2026-09": "12345"},
            "actuals": {"2026-09": "0"},
        }
    ]
    pdf = render_subsidiary_pdf(
        snapshot,
        {"status": "UNVERIFIED", "snapshot_sha256": "a" * 64},
        ROOT / "resources/fonts/ReportingSerif.ttf",
    )
    assert pdf.startswith(b"%PDF")
    assert b"SYNTHETIC-12345" in pdf
    assert b"TEST-CODE-A" in pdf
    (tmp_path / "head-reference.pdf").write_bytes(pdf)
