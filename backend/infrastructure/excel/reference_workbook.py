"""Read the supplied monthly/weekly forms without inventing daily coordinates."""

from __future__ import annotations

import ast
import io
import re
import zipfile
from datetime import date, datetime
from decimal import Decimal
from typing import Any, cast

from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries


class FormulaReader:
    def __init__(self, cells: dict[str, dict[str, Any]]) -> None:
        self.cells = cells
        self.cache: dict[str, Any] = {}
        self.active: set[str] = set()

    def cell(self, address: str) -> Any:
        address = address.replace("$", "")
        if address in self.cache:
            return self.cache[address]
        if address in self.active or len(self.active) > 100:
            raise ValueError("Циклическая или слишком глубокая формула")
        item = self.cells.get(address, {})
        value = item.get("value")
        if item.get("kind") == "f":
            self.active.add(address)
            try:
                value = self.formula(str(value))
            finally:
                self.active.remove(address)
        elif item.get("kind") == "n":
            value = Decimal(str(value))
        self.cache[address] = value
        return value

    def formula(self, formula: str) -> Any:
        if len(formula) > 10000:
            raise ValueError("Слишком длинная формула")
        expression = formula.removeprefix("=").strip()
        if "!" in expression or "[" in expression or "]" in expression:
            raise ValueError("Внешние и межлистовые ссылки требуют проверки")
        # Tokenize references only outside quoted Excel strings.
        parts = re.split(r'("(?:[^"]|"")*")', expression)

        def reference(match: re.Match[str]) -> str:
            return 'REF("' + match.group(0) + '")'

        for index in range(0, len(parts), 2):
            parts[index] = re.sub(
                r"(?<![A-Za-z0-9_])\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?",
                reference,
                parts[index],
            )
        tree = ast.parse("".join(parts), mode="eval")
        if sum(1 for _ in ast.walk(tree)) > 6000:
            raise ValueError("Слишком сложная формула")
        return self.node(tree.body)

    def node(self, node: ast.AST) -> Any:
        if isinstance(node, ast.Constant) and isinstance(node.value, (str, int, float)):
            return node.value if isinstance(node.value, str) else Decimal(str(node.value))
        if isinstance(node, ast.UnaryOp) and isinstance(node.op, (ast.UAdd, ast.USub)):
            value = self.number(self.node(node.operand))
            return -value if isinstance(node.op, ast.USub) else value
        if isinstance(node, ast.BinOp):
            left, right = self.number(self.node(node.left)), self.number(self.node(node.right))
            if isinstance(node.op, ast.Add):
                return left + right
            if isinstance(node.op, ast.Sub):
                return left - right
            if isinstance(node.op, ast.Mult):
                return left * right
            if isinstance(node.op, ast.Div):
                return left / right
        if isinstance(node, ast.Call) and isinstance(node.func, ast.Name) and not node.keywords:
            name = node.func.id.upper()
            args = [self.node(arg) for arg in node.args]
            if name == "REF" and len(args) == 1 and isinstance(args[0], str):
                ref = args[0]
                if ":" not in ref:
                    return self.cell(ref)
                from openpyxl.utils import get_column_letter

                x1, y1, x2, y2 = cast(
                    tuple[int, int, int, int], range_boundaries(ref.replace("$", ""))
                )
                if (x2 - x1 + 1) * (y2 - y1 + 1) > 20000:
                    raise ValueError("Слишком большой диапазон")
                return [
                    self.cell(f"{get_column_letter(x)}{y}")
                    for y in range(y1, y2 + 1)
                    for x in range(x1, x2 + 1)
                ]
            flat = [v for arg in args for v in (arg if isinstance(arg, list) else [arg])]
            if name == "SUM":
                return sum((v for v in flat if isinstance(v, Decimal)), Decimal(0))
            if name == "CONCATENATE":
                return "".join("" if v is None else str(v) for v in flat)
            if name == "CHAR" and len(args) == 1:
                return chr(int(self.number(args[0])))
        raise ValueError("Неподдерживаемая формула в образце")

    @staticmethod
    def number(value: Any) -> Decimal:
        if value is None:
            return Decimal(0)
        if not isinstance(value, Decimal):
            raise ValueError("В формуле вместо числа текст")
        return value


def calculate(document: dict[str, Any]) -> dict[str, Any]:
    errors = []
    for sheet in document["sheets"]:
        reader = FormulaReader(sheet["cells"])
        for address, cell in sheet["cells"].items():
            try:
                value = reader.cell(address)
                cell["display"] = "" if value is None else str(value)
                if cell.get("kind") == "f":
                    cached = cell.get("cached_value")
                    cell["formula_status"] = (
                        "CALCULATED_NO_CACHE"
                        if cached is None
                        else "CACHE_VERIFIED"
                        if str(cached) == str(value)
                        else "CACHE_DIFFERS"
                    )
                cell.pop("error", None)
            except (ValueError, SyntaxError, ArithmeticError, TypeError, RecursionError) as exc:
                cell["display"] = "#ОШИБКА"
                cell["error"] = str(exc)
                cell["formula_status"] = "UNSUPPORTED"
                errors.append(f"{sheet['name']}!{address}: {exc}")
    document["errors"] = errors
    return document


def read_reference(content: bytes) -> dict[str, Any]:
    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        if sum(x.file_size for x in archive.infolist()) > 100 * 1024 * 1024:
            raise ValueError("Распакованная книга превышает 100 МБ")
    workbook = load_workbook(io.BytesIO(content), data_only=False, keep_links=False)
    cached_workbook = load_workbook(io.BytesIO(content), data_only=True, keep_links=False)
    try:
        sheets = []
        for sheet in workbook:
            if sheet.max_row * sheet.max_column > 500_000:
                raise ValueError("Лист превышает 500 000 ячеек")
            cells = {}
            for row in sheet:
                for c in row:
                    if c.value is None:
                        continue
                    value = c.value
                    kind = (
                        "f"
                        if c.data_type == "f"
                        else (
                            "n"
                            if isinstance(value, (int, float)) and not isinstance(value, bool)
                            else "s"
                        )
                    )
                    if isinstance(value, (date, datetime)):
                        value, kind = value.strftime("%Y-%m-%d"), "d"
                    cells[c.coordinate] = {
                        "value": str(value),
                        "kind": kind,
                        "number_format": c.number_format,
                        "cached_value": (
                            str(cached_workbook[sheet.title][c.coordinate].value)
                            if kind == "f"
                            and cached_workbook[sheet.title][c.coordinate].value is not None
                            else None
                        ),
                    }
            sheets.append(
                {
                    "name": sheet.title,
                    "rows": sheet.max_row,
                    "columns": sheet.max_column,
                    "cells": cells,
                    "merges": [str(m) for m in sheet.merged_cells.ranges],
                    "state": sheet.sheet_state,
                    "hidden_rows": [i for i, d in sheet.row_dimensions.items() if d.hidden],
                    "hidden_columns": [i for i, d in sheet.column_dimensions.items() if d.hidden],
                }
            )
        document = calculate({"sheets": sheets, "warnings": []})
        # One server recognizer is authoritative for role classification and UI sources.
        from backend.application.import_recognition import recognize_document

        document["recognition"] = recognize_document(document)
        kinds = {s.get("kind") for s in sheets if s.get("kind")}
        document["report_types"] = sorted(kinds)
        document["report_type"] = next((s["kind"] for s in sheets if s.get("kind")), "")
        if not kinds:
            raise ValueError("Не найдены заголовки отчётной таблицы: обозначение и входимость")
        return document
    finally:
        workbook.close()
        cached_workbook.close()


def export_reference(content: bytes, document: dict[str, Any]) -> bytes:
    workbook = load_workbook(io.BytesIO(content), keep_links=False)
    try:
        for sheet in document["sheets"]:
            for address, item in sheet["cells"].items():
                if item.get("edited"):
                    value = item["value"]
                    workbook[sheet["name"]][address] = (
                        Decimal(value) if item["kind"] == "n" else value
                    )
                    if item["kind"] == "s":
                        workbook[sheet["name"]][address].data_type = "s"
        output = io.BytesIO()
        workbook.save(output)
        return output.getvalue()
    finally:
        workbook.close()
