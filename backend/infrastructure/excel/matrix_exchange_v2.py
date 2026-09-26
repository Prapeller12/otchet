"""Versioned matrix exchange metadata; addresses are never entity identities."""

from __future__ import annotations

import calendar
import hashlib
import json
from collections.abc import Mapping
from datetime import date
from pathlib import Path
from typing import Any
from uuid import NAMESPACE_URL, uuid5

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Protection
from openpyxl.utils import get_column_letter

from backend.application.excel_reports import ExcelWorkbookValidationError, WorkbookIssue
from backend.application.report_calendar import reporting_weeks

METADATA_SHEET = "_Обмен v2"
CLEAR_TOKEN = "#CLEAR"
CALENDAR = "MONDAY_SUNDAY_SHORT_FRAGMENTS_MERGED_V1"


def canonical(value: object) -> str:
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def periods(matrix: Mapping[str, Any]) -> list[dict[str, Any]]:
    result = []
    for column in matrix["time_columns"]:
        item = dict(column)
        group = str(column.get("group_label", ""))
        try:
            year, month = map(int, group.split("-"))
            start = date(year, month, 1)
            end = date(year, month, calendar.monthrange(year, month)[1])
            item.update(start=start.isoformat(), end=end.isoformat(), calendar=CALENDAR)
            if column.get("kind") in {"USED", "DELIVERED", "SUPPLIED"}:
                week = next(
                    (
                        w
                        for w in reporting_weeks(year, month)
                        if w.start.isoformat() == str(column["id"])[:10]
                    ),
                    None,
                )
                if week:
                    item.update(start=week.start.isoformat(), end=week.end.isoformat())
                else:
                    item.update(start=None, end=None, calendar="UNRESOLVED")
                if column.get("period_start") and column.get("period_end"):
                    exact_start = date.fromisoformat(str(column["period_start"]))
                    exact_end = date.fromisoformat(str(column["period_end"]))
                    if exact_start > exact_end:
                        raise ValueError("Reversed weekly period")
                    item.update(start=exact_start.isoformat(), end=exact_end.isoformat())
            elif len(str(column.get("id", ""))) == 10:
                day = date.fromisoformat(column["id"])
                item.update(start=day.isoformat(), end=day.isoformat(), calendar="DAILY")
        except (ValueError, TypeError):
            item.update(start=None, end=None, calendar="UNRESOLVED")
        result.append(item)
    return result


def _literal(sheet: Any, row: int, column: int, value: object) -> None:
    cell = sheet.cell(row, column, str(value))
    cell.data_type = "s"


def write_metadata(workbook: Workbook, matrix: Mapping[str, Any]) -> None:
    """Keep metadata in small JSON records (Excel cells are limited to 32767 chars)."""
    sheet = workbook["Отчёт"]
    meta = workbook.create_sheet(METADATA_SHEET)
    meta.sheet_state = "veryHidden"
    identity = dict(matrix.get("exchange_identity", {}))
    snapshot = {
        "schema_version": 2,
        "exchange_identity": identity,
        "year": matrix.get("year"),
        "report_type": matrix["report_type"],
        "organization_id": matrix["organization_id"],
        "presentation": matrix.get("presentation", {}),
        "structure": matrix.get("exchange_structure", []),
        "rows": [
            {
                k: v
                for k, v in row.items()
                if k
                in {"id", "workspace_id", "group_id", "supplier_id", "left_values", "metric_code"}
            }
            for row in matrix["rows"]
        ],
        "coordinates": [
            {"source_cell": row[0], "coordinate": json.loads(row[1])}
            for row in workbook["_Системная карта"].iter_rows(
                min_row=6, max_col=3, values_only=True
            )
            if isinstance(row[0], str) and isinstance(row[1], str)
        ],
        "periods": periods(matrix),
        "non_imported_fields": [
            "Расчётные остатки, дефицит и проценты: пересчитываются программой",
            "Вычисляемые годовые итоги: пересчитываются из первичных месячных значений",
        ],
    }
    payload = canonical(snapshot)
    meta.append(["REPORTING_SYSTEM_EXCHANGE", 2, hashlib.sha256(payload.encode()).hexdigest()])
    for offset in range(0, len(payload), 30000):
        _literal(meta, meta.max_row + 1, 1, payload[offset : offset + 30000])
    # A row token moves with an inserted/moved Excel row; stale coordinate maps fail closed.
    token_column = sheet.max_column + 1
    token_letter = get_column_letter(token_column)
    sheet.column_dimensions[token_letter].hidden = True
    sheet.cell(6, token_column, "_exchange_row_identity")
    guards: list[dict[str, Any]] = []
    for index, row in enumerate(snapshot["rows"], 7):
        token = str(uuid5(NAMESPACE_URL, canonical([identity, row.get("id"), index])))
        sheet.cell(index, token_column, token)
        guards.append({"cell": f"{token_letter}{index}", "value": token})
    # Header cells and merges guard column insertion/reordering, independent of visibility.
    for row in sheet.iter_rows(min_row=5, max_row=6, max_col=token_column):
        for cell in row:
            if cell.value is not None:
                guards.append({"cell": cell.coordinate, "value": cell.value})
    editable_fields = []
    for index, _row in enumerate(snapshot["rows"], 7):
        for column_index, column in enumerate(matrix["left_columns"], 1):
            cell = sheet.cell(index, column_index)
            # Shared/merged cells have one authoritative source, not repeated supplier copies.
            if cell.__class__.__name__ == "MergedCell":
                continue
            editable_fields.append(
                {
                    "sheet": "Отчёт",
                    "cell": cell.coordinate,
                    "path": ["rows", index - 7, "left_values", column["id"]],
                    "value": cell.value,
                    "type": "text",
                }
            )
    if "Шапка и выпуск" in workbook:
        header = workbook["Шапка и выпуск"]
        for index, key in enumerate(("product_designation", "product_name", "factory_name"), 1):
            header.cell(index, 2).protection = Protection(locked=False)
            editable_fields.append(
                {
                    "sheet": header.title,
                    "cell": f"B{index}",
                    "path": ["presentation", "header", key],
                    "value": header.cell(index, 2).value,
                    "type": "text",
                }
            )
        guards.append({"sheet": header.title, "cell": "B4", "value": header["B4"].value})
    if "Месячные планы" in workbook:
        plans = workbook["Месячные планы"]
        for index in range(2, 14):
            period = plans.cell(index, 1).value
            if not isinstance(period, str) or len(period) != 7:
                continue
            guards.append({"sheet": plans.title, "cell": f"A{index}", "value": period})
            for column_index, key in ((2, "plans"), (3, "actuals")):
                cell = plans.cell(index, column_index)
                cell.protection = Protection(locked=False)
                editable_fields.append(
                    {
                        "sheet": plans.title,
                        "cell": cell.coordinate,
                        "path": ["presentation", key, period],
                        "value": cell.value,
                        "type": "quantity",
                    }
                )
    controls = canonical(
        {
            "guards": guards,
            "fields": editable_fields,
            "merges": sorted(str(r) for r in sheet.merged_cells.ranges),
        }
    )
    meta.cell(meta.max_row + 1, 1, "CONTROLS")
    for offset in range(0, len(controls), 30000):
        _literal(meta, meta.max_row + 1, 1, controls[offset : offset + 30000])


def read_metadata(workbook: Workbook) -> tuple[dict[str, Any], dict[str, Any]]:
    if METADATA_SHEET not in workbook:
        raise ExcelWorkbookValidationError("В книге v2 отсутствуют метаданные обмена")
    meta = workbook[METADATA_SHEET]
    if meta["A1"].value != "REPORTING_SYSTEM_EXCHANGE" or meta["B1"].value != 2:
        raise ExcelWorkbookValidationError("Повреждена версия метаданных обмена")
    parts: list[str] = []
    controls: list[str] = []
    target = parts
    for row in meta.iter_rows(min_row=2, max_col=1, values_only=True):
        if row[0] == "CONTROLS":
            target = controls
        elif isinstance(row[0], str):
            target.append(row[0])
        else:
            raise ExcelWorkbookValidationError("Повреждены метаданные обмена")
    payload = "".join(parts)
    if hashlib.sha256(payload.encode()).hexdigest() != meta["C1"].value:
        raise ExcelWorkbookValidationError("Контрольная сумма метаданных обмена не совпала")
    try:
        snapshot, control = json.loads(payload), json.loads("".join(controls))
        if not isinstance(snapshot, dict) or not isinstance(control, dict):
            raise ValueError("metadata must be objects")
        return snapshot, control
    except (ValueError, TypeError) as exc:
        raise ExcelWorkbookValidationError("Повреждены метаданные обмена") from exc


def read_exchange_snapshot(source: Path) -> dict[str, Any] | None:
    """Bridge preflight: structure can be staged before mapping incoming coordinates."""
    workbook = load_workbook(source, keep_links=False)
    try:
        if METADATA_SHEET not in workbook:
            return None
        snapshot = read_metadata(workbook)[0]
        snapshot["coordinates"] = [
            {"source_cell": row[0], "coordinate": json.loads(row[1])}
            for row in workbook["_Системная карта"].iter_rows(
                min_row=6, max_col=3, values_only=True
            )
            if isinstance(row[0], str) and isinstance(row[1], str)
        ]
        return snapshot
    finally:
        workbook.close()


def validate_metadata(
    workbook: Workbook, matrix: Mapping[str, Any]
) -> tuple[dict[str, Any], list[WorkbookIssue]]:
    snapshot, controls = read_metadata(workbook)
    issues = []
    source_dataset = snapshot.get("exchange_identity", {}).get("dataset_id")
    target_dataset = matrix.get("exchange_identity", {}).get("dataset_id")
    accepted_source = matrix.get("exchange_source_dataset_id")
    if source_dataset != target_dataset and not (
        source_dataset and accepted_source == source_dataset
    ):
        issues.append(
            WorkbookIssue(
                None,
                "FOREIGN_DATASET",
                "Файл создан в другой базе. Требуется сопоставление структуры и идентификаторов.",
            )
        )
    if snapshot.get("year") != matrix.get("year"):
        issues.append(WorkbookIssue(None, "YEAR_MISMATCH", "Год книги отличается от года отчёта"))
    for guard in controls.get("guards", []):
        name = guard.get("sheet", "Отчёт")
        if name not in workbook or workbook[name][guard["cell"]].value != guard["value"]:
            issues.append(
                WorkbookIssue(
                    f"{name}!{guard['cell']}",
                    "LAYOUT_CHANGED",
                    "Строки или колонки перемещены; системная карта устарела. Повторите экспорт.",
                )
            )
    if sorted(str(r) for r in workbook["Отчёт"].merged_cells.ranges) != controls.get("merges"):
        issues.append(
            WorkbookIssue(None, "LAYOUT_CHANGED", "Изменена структура объединённых ячеек")
        )
    changes = []
    for field in controls.get("fields", []):
        name, address = field["sheet"], field["cell"]
        if name not in workbook:
            issues.append(WorkbookIssue(name, "MISSING_SHEET", "Отсутствует лист обмена"))
            continue
        value = workbook[name][address].value
        if value == field["value"] or value is None or value == "":
            continue
        changes.append(
            {**field, "source_cell": f"{name}!{address}", "before": field["value"], "after": value}
        )
    return {
        "schema_version": 2,
        "snapshot": snapshot,
        "structural_changes": changes,
        "non_imported_fields": snapshot.get("non_imported_fields", []),
    }, issues


def review_sheets(
    workbook: Workbook, known_sheets: set[str], raw_decisions: object
) -> tuple[list[dict[str, object]], list[WorkbookIssue]]:
    """Unknown sheets stay visible in the review; exclusions always carry a reason."""
    issues: list[WorkbookIssue] = []
    sheets: list[dict[str, object]] = []
    decisions: Mapping[str, Any] = {}
    if raw_decisions is not None:
        if isinstance(raw_decisions, Mapping):
            decisions = raw_decisions
        else:
            issues.append(
                WorkbookIssue(
                    None, "SHEET_DECISION_INVALID", "Решения по листам должны быть объектом"
                )
            )
    for name in decisions:
        if name not in workbook.sheetnames:
            issues.append(
                WorkbookIssue(str(name), "UNKNOWN_SHEET", "Решение ссылается на отсутствующий лист")
            )
    for name in workbook.sheetnames:
        known = name in known_sheets
        entry: dict[str, object] = {
            "name": name,
            "state": workbook[name].sheet_state,
            "known": known,
            "requires_decision": not known,
            "included": True if known else None,
            "reason": "",
        }
        sheets.append(entry)
        decision = decisions.get(name)
        if known:
            if isinstance(decision, Mapping) and decision.get("include") is False:
                issues.append(
                    WorkbookIssue(
                        name,
                        "SYSTEM_SHEET_REQUIRED",
                        "Рабочий или служебный лист собственного формата исключить нельзя",
                    )
                )
            continue
        if decision is None:
            issues.append(
                WorkbookIssue(
                    name,
                    "SHEET_REVIEW_REQUIRED",
                    "Неизвестный лист: выберите включение или исключение с причиной",
                )
            )
            continue
        if (
            not isinstance(decision, Mapping)
            or type(decision.get("include")) is not bool
            or not isinstance(decision.get("reason", ""), str)
            or set(decision) - {"include", "reason"}
        ):
            issues.append(
                WorkbookIssue(
                    name,
                    "SHEET_DECISION_INVALID",
                    "Решение по листу должно содержать включение и текстовую причину",
                )
            )
            continue
        include = decision["include"]
        reason = decision.get("reason", "").strip()
        entry.update(included=include, reason=reason)
        if include:
            issues.append(
                WorkbookIssue(
                    name,
                    "UNSUPPORTED_SHEET",
                    "Дополнительный лист не входит в системную карту. "
                    "Сохраните его отдельной книгой "
                    "и импортируйте как внешний отчёт с проверкой структуры и периодов.",
                )
            )
        elif not reason:
            issues.append(
                WorkbookIssue(
                    name,
                    "SHEET_EXCLUSION_REASON_REQUIRED",
                    "Укажите причину исключения листа; без неё лист не будет пропущен",
                )
            )
    return sheets, issues
