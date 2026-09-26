"""Versioned, offline recognition of enterprise workbook fields and exact periods.

No worksheet ordinal is a business identifier; ambiguous dates stay unresolved.
"""

from __future__ import annotations

import calendar
import hashlib
import json
import re
from datetime import date
from decimal import Decimal, InvalidOperation
from functools import lru_cache
from pathlib import Path
from typing import Any

from openpyxl.utils.cell import coordinate_to_tuple, get_column_letter, range_boundaries


@lru_cache(maxsize=1)
def profile() -> dict[str, Any]:
    root = Path(__file__).resolve().parents[2]
    path = root / "config/import_profiles/enterprise.v1.json"
    if not path.exists():  # portable backend lives under app/backend
        path = root.parent / "config/import_profiles/enterprise.v1.json"
    return dict(json.loads(path.read_text(encoding="utf-8")))


def normalize(value: object) -> str:
    return " ".join(str(value or "").lower().replace("ё", "е").split())


def decimal_text(value: object) -> str:
    raw = str(value).strip().replace("\u00a0", "").replace("\u202f", "").replace(" ", "")
    if not re.fullmatch(r"[+-]?\d+(?:[.,]\d+)?", raw):
        raise ValueError("Укажите число; текст, даты и прочерки не являются количеством")
    try:
        number = Decimal(raw.replace(",", "."))
    except InvalidOperation as exc:
        raise ValueError("Некорректное число") from exc
    if not number.is_finite() or abs(number) > Decimal("1e50"):
        raise ValueError("Значение вне допустимого диапазона")
    return format(number, "f")


def merged_cells(sheet: dict[str, Any]) -> dict[str, str]:
    result = {}
    for region in sheet.get("merges", []):
        x1, y1, x2, y2 = range_boundaries(region)
        if None in (x1, y1, x2, y2):
            continue
        assert x1 is not None and y1 is not None and x2 is not None and y2 is not None
        anchor = f"{get_column_letter(x1)}{y1}"
        for y in range(y1, y2 + 1):
            for x in range(x1, x2 + 1):
                result[f"{get_column_letter(x)}{y}"] = anchor
    return result


def recognize_document(document: dict[str, Any]) -> dict[str, Any]:
    spec = profile()
    sources: dict[str, Any] = {}
    fields: dict[str, Any] = {}
    actions: list[dict[str, Any]] = []
    issues: list[dict[str, str]] = []
    period_blocks: dict[str, Any] = {}
    structure: list[Any] = []
    for index, sheet in enumerate(document["sheets"]):
        cells = sheet["cells"]
        merged = merged_cells(sheet)

        def cell_at(row: int, col: int) -> dict[str, Any]:
            address = f"{get_column_letter(col)}{row}"
            return dict(cells.get(merged.get(address, address), {}))  # noqa: B023

        def text_at(row: int, col: int) -> str:
            item = cell_at(row, col)
            return str(item.get("display", item.get("value", "")))

        header_rows = []
        for row in range(1, sheet["rows"] + 1):
            if any(
                "входимость"
                in normalize(cells.get(f"{get_column_letter(col)}{row}", {}).get("value"))
                for col in range(1, sheet["columns"] + 1)
            ):
                header_rows.append(row)
        sheet["header_rows"] = header_rows
        if not header_rows:
            issues.append(
                {
                    "source_cell": f"{index}:",
                    "code": "SHEET_REQUIRED",
                    "message": f"Лист «{sheet['name']}» не распознан; укажите причину исключения",
                }
            )
        if sheet.get("state", "visible") != "visible":
            issues.append(
                {
                    "source_cell": f"{index}:",
                    "code": "SHEET_REQUIRED",
                    "message": f"Скрытый лист «{sheet['name']}»: "
                    "подтвердите включение или исключение",
                }
            )
        if not header_rows:
            for address, item in cells.items():
                item["role"] = "REQUISITE"
                fields[f"{index}:{address}"] = {
                    "sheet": sheet["name"],
                    "address": address,
                    "role": "REQUISITE",
                    "value": item["value"],
                }
            continue
        header = header_rows[0]
        headers = {col: normalize(text_at(header, col)) for col in range(1, sheet["columns"] + 1)}
        structural_cols: dict[str, int] = {}
        for key, names in spec["structure_headers"].items():
            for col, label in headers.items():
                if any(normalize(name) in label for name in names):
                    structural_cols.setdefault(key, col)
        # This profile is selected by semantic header signatures, not fixed row numbers.
        kind = "HEAD_SITE" if any(h == "план" for h in headers.values()) else "SUBSIDIARY"
        sheet["kind"] = kind
        sheet["structure_columns"] = structural_cols
        structure.append([sheet["name"], headers, sheet.get("merges", [])])
        title_text = " ".join(
            str(v.get("value", ""))
            for a, v in cells.items()
            if coordinate_to_tuple(a)[0] < header and coordinate_to_tuple(a)[1] < 4
        )
        title_years = set(re.findall(r"\b20\d{2}\b", title_text))
        default_year = int(next(iter(title_years))) if len(title_years) == 1 else None
        product_context = next(
            (
                str(v.get("value", ""))
                for a, v in cells.items()
                if coordinate_to_tuple(a)[0] < header
                and "\n" in str(v.get("value", ""))
                and re.search(r"[A-ZА-Я]+\.\d", str(v.get("value", "")))
            ),
            "",
        )
        # Product-code plans above the component table are independent primary values.
        presentation_addresses: set[str] = set()
        monthly_columns = {
            col: label for col, label in headers.items() if label in {"план", "факт"}
        }
        production_codes = []
        if kind == "HEAD_SITE" and monthly_columns:
            first_month_col = min(monthly_columns)
            for source_row in range(1, header):
                label_cells = [
                    (col, text_at(source_row, col).strip())
                    for col in range(1, first_month_col)
                    if cell_at(source_row, col).get("kind") == "s"
                    and text_at(source_row, col).strip()
                ]
                if not label_cells:
                    continue
                label_col, product_label = label_cells[-1]
                if label_col <= max(structural_cols.values(), default=0):
                    continue
                code: dict[str, Any] = {
                    "id": "X" + hashlib.sha256(product_label.encode()).hexdigest()[:20].upper(),
                    "label": product_label,
                    "plans": {},
                    "actuals": {},
                }
                for col, metric_label in monthly_columns.items():
                    address = f"{get_column_letter(col)}{source_row}"
                    item = cells.get(address)
                    if not item or item.get("kind") not in {"n", "f"}:
                        continue
                    period, _ = parse_period(
                        [text_at(r, col) for r in range(1, source_row)], default_year, spec
                    )
                    if not period:
                        issues.append(
                            {
                                "source_cell": f"{index}:{address}",
                                "code": "PERIOD_REQUIRED",
                                "message": "Не определён месяц плана/выпуска в шапке",
                            }
                        )
                        continue
                    try:
                        value = decimal_text(item.get("display", item.get("value", "")))
                    except ValueError as exc:
                        issues.append(
                            {
                                "source_cell": f"{index}:{address}",
                                "code": "FORMULA_ERROR",
                                "message": str(exc),
                            }
                        )
                        continue
                    code["plans" if metric_label == "план" else "actuals"][period["start"][:7]] = (
                        value
                    )
                    presentation_addresses.add(address)
                production_codes.append(code)
            if production_codes:
                actions.append(
                    {
                        "kind": "UPSERT_PRESENTATION",
                        "source_key": f"{index}:header",
                        "sheet_index": index,
                        "report_type": kind,
                        "patch": {"production_codes": production_codes},
                        "errors": [],
                    }
                )
        section_rows: set[int] = set()
        positions: dict[int, dict[str, str]] = {}
        for row in range(header + 1, sheet["rows"] + 1):
            if row in header_rows:
                continue
            anchors = [
                merged.get(f"{get_column_letter(structural_cols[key])}{row}")
                for key in ("code", "name", "norm")
                if key in structural_cols
            ]
            if len(anchors) == 3 and anchors[0] and len(set(anchors)) == 1:
                section_rows.add(row)
                continue
            position = {key: text_at(row, col).strip() for key, col in structural_cols.items()}
            if not position.get("code"):
                continue
            if normalize(position["code"]) in {"обозначение", "обозначние"}:
                continue
            for key in ("code", "name", "structure_number", "norm", "manufacturer", "contract"):
                position.setdefault(key, "")
            position["parent_code"] = ""
            positions[row] = position
        structure.append(
            {
                "sheet": sheet["name"],
                "positions": positions,
                "period_headers": {
                    a: c.get("value")
                    for a, c in cells.items()
                    if coordinate_to_tuple(a)[0] < header and c.get("kind") in {"s", "d"}
                },
            }
        )
        number_codes: dict[str, set[str]] = {}
        for position in positions.values():
            number_codes.setdefault(position["structure_number"], set()).add(position["code"])
        action_lookup: dict[tuple[str, str], dict[str, Any]] = {}
        for row, position in positions.items():
            number = position["structure_number"]
            parent_number = number.rsplit(".", 1)[0] if "." in number else ""
            parent_codes = number_codes.get(parent_number, set())
            if len(parent_codes) == 1:
                position["parent_code"] = next(iter(parent_codes))
            key = (position["code"], position["parent_code"])
            action = action_lookup.get(key)
            if action is None:
                action = {
                    "kind": "UPSERT_POSITION",
                    "source_key": f"{index}:{row}",
                    "sheet_index": index,
                    "report_type": kind,
                    "template_group_id": spec["templates"].get(kind),
                    "position": dict(position),
                    "product_context": product_context,
                    "organization_context": sheet["name"],
                    "source_rows": [],
                    "suppliers": [],
                    "errors": [],
                    "weekly_supply": kind == "SUBSIDIARY"
                    and any(
                        "недел" in normalize(v.get("value"))
                        for a, v in cells.items()
                        if coordinate_to_tuple(a)[0] < header
                    ),
                }
                action_lookup[key] = action
                actions.append(action)
            for field in ("norm", "contract"):
                if position[field]:
                    try:
                        position[field] = decimal_text(position[field])
                        if action["position"].get(field):
                            action["position"][field] = decimal_text(action["position"][field])
                    except ValueError as exc:
                        action["errors"].append(
                            {
                                "code": "STRUCTURE_REQUIRED",
                                "message": f"{sheet['name']} строка {row}: {field}: {exc}",
                            }
                        )
            if action["position"]["name"] != position["name"]:
                action["errors"].append(
                    {
                        "code": "IDENTITY_CONFLICT",
                        "message": "Одно обозначение имеет разные названия; "
                        "уточните идентичность позиции",
                    }
                )
            action["source_rows"].append(row)
            if position["manufacturer"]:
                supplier = {
                    "name": position["manufacturer"],
                    "contract": position["contract"],
                    "source_row": row,
                }
                if not any(s["name"] == supplier["name"] for s in action["suppliers"]):
                    action["suppliers"].append(supplier)
        for action in action_lookup.values():
            if not action["suppliers"]:
                action["errors"].append(
                    {
                        "code": "MANUFACTURER_REQUIRED",
                        "message": f"{sheet['name']} строки {action['source_rows']}: "
                        "укажите изготовителя "
                        f"для позиции {action['position']['code']}",
                    }
                )
        for address, item in cells.items():
            row, col = coordinate_to_tuple(address)
            raw = str(item.get("display", item.get("value", "")))
            label = headers.get(col, "")
            role = "REQUISITE"
            metric = None
            unit = "шт."
            position = positions.get(row, {})
            period = None
            errors = []
            if "%" in item.get("number_format", ""):
                role, metric, unit = "CONTROL", "COMPLETION", "%"
            elif row in section_rows:
                role = "DECORATION"
            elif address in presentation_addresses:
                role = "STRUCTURE"
            elif row <= header or row in header_rows:
                role = "REQUISITE"
            elif col in structural_cols.values():
                role = "STRUCTURE"
            elif position:
                role = "FORMULA" if item.get("kind") == "f" else "INPUT"
                if any(term in label for term in spec["control_headers"]):
                    role = "CONTROL"
                for metric_code, labels in spec["metrics"].items():
                    if any(normalize(term) in label for term in labels):
                        metric = metric_code
                        break
                if metric == "FACT":
                    metric = "SUPPLIED" if kind == "SUBSIDIARY" else "PRODUCED"
                block_key = f"{index}:{get_column_letter(col)}"
                if block_key not in period_blocks:
                    above = [text_at(r, col) for r in range(1, header)]
                    period, block = parse_period(above, default_year, spec)
                    block.update(key=block_key, sheet=sheet["name"], address=get_column_letter(col))
                    if metric and role != "CONTROL":
                        period_blocks[block_key] = block
                    if block.get("year") and default_year and block["year"] != default_year:
                        message = "Год колонок отличается от года заголовка; сохранены даты колонок"
                        if message not in document["warnings"]:
                            document["warnings"].append(message)
                else:
                    period = period_blocks[block_key].get("period")
                if role in {"INPUT", "FORMULA"} and raw.strip():
                    try:
                        raw = decimal_text(raw)
                    except ValueError as exc:
                        errors.append({"code": "QUANTITY_REQUIRED", "message": str(exc)})
                    if metric is None:
                        errors.append(
                            {"code": "MAPPING_REQUIRED", "message": "Уточните смысл показателя"}
                        )
                    if period is None:
                        errors.append(
                            {
                                "code": "PERIOD_REQUIRED",
                                "message": "Укажите точные даты периода для колонки",
                            }
                        )
            item["role"] = role
            item["unit"] = unit
            entry = {
                "sheet": sheet["name"],
                "address": address,
                "role": role,
                "source_label": label,
                "unit": unit,
                "value": raw,
                "raw_value": item.get("value"),
                "number_format": item.get("number_format", "General"),
                "formula": str(item["value"]) if item.get("kind") == "f" else "",
                "formula_status": item.get("formula_status", "NOT_FORMULA"),
                "position": dict(position),
                "product_context": product_context,
                "metric": metric,
                "period": period,
                "errors": errors,
                "period_block": f"{index}:{get_column_letter(col)}",
            }
            if item.get("error") and role not in {"INPUT", "FORMULA"}:
                issues.append(
                    {
                        "source_cell": f"{index}:{address}",
                        "code": "FORMULA_ERROR",
                        "message": str(item["error"]),
                    }
                )
            if item.get("error"):
                entry["errors"].append({"code": "FORMULA_ERROR", "message": item["error"]})
            fields[f"{index}:{address}"] = entry
            if role in {"INPUT", "FORMULA"} and raw.strip():
                sources[f"{index}:{address}"] = entry
    return {
        "profile_id": spec["id"],
        "profile_version": spec["version"],
        "structure_fingerprint": hashlib.sha256(
            json.dumps(structure, sort_keys=True, ensure_ascii=False).encode()
        ).hexdigest(),
        "sources": sources,
        "fields": fields,
        "structural_actions": actions,
        "period_blocks": list(period_blocks.values()),
        "issues": issues,
        "sheets": [
            {
                "index": i,
                "name": s["name"],
                "state": s.get("state", "visible"),
                "decision_required": not s.get("kind") or s.get("state", "visible") != "visible",
            }
            for i, s in enumerate(document["sheets"])
        ],
    }


def parse_period(
    above: list[str], default_year: int | None, spec: dict[str, Any]
) -> tuple[dict[str, str] | None, dict[str, Any]]:
    month = None
    year = default_year
    week = ""
    explicit = None
    for text in above:
        clean = normalize(text)
        if "недел" in clean:
            week = text
        if re.fullmatch(r"20\d{2}-\d{2}-\d{2}", clean):
            parsed = date.fromisoformat(clean)
            year, month = parsed.year, parsed.month
            continue
        for i, name in enumerate(spec["months"], 1):
            if re.match(r"^" + name, clean):
                month = i
                found = re.search(r"\b20\d{2}\b", clean)
                if found:
                    year = int(found[0])
        match = re.search(r"\b(\d{1,2})[–—-](\d{1,2})\b", clean)
        if match:
            explicit = (int(match[1]), int(match[2]))
    period = None
    if year and month and (not week or explicit):
        days = explicit or (1, calendar.monthrange(year, month)[1])
        try:
            start, end = date(year, month, days[0]), date(year, month, days[1])
            if start <= end:
                period = {
                    "start": start.isoformat(),
                    "end": end.isoformat(),
                    "calendar": "EXPLICIT" if explicit else "MONTH",
                }
        except ValueError:
            pass
    return period, {
        "label": " / ".join(t for t in above if t.strip()),
        "year": year,
        "month": month,
        "week_label": week,
        "period": period,
        "errors": []
        if period
        else [{"code": "PERIOD_REQUIRED", "message": "В источнике нет точных границ периода"}],
    }


def apply_structure_overrides(recognition: dict[str, Any], overrides: object) -> dict[str, Any]:
    """Apply explicit, auditable identity corrections to the shared recognition result."""
    import copy

    if overrides is None:
        return recognition
    if not isinstance(overrides, dict):
        raise ValueError("Исправления структуры должны быть объектом")
    result = copy.deepcopy(recognition)
    actions = {
        a["source_key"]: a for a in result["structural_actions"] if a["kind"] == "UPSERT_POSITION"
    }
    for key, patch in overrides.items():
        if key not in actions or not isinstance(patch, dict):
            raise ValueError("Неизвестная позиция для исправления структуры")
        allowed = {"code", "name", "manufacturer", "norm", "parent_code", "reason"}
        if (
            patch.keys() - allowed
            or not isinstance(patch.get("reason"), str)
            or not patch["reason"].strip()
        ):
            raise ValueError("Укажите причину исправления структуры")
        updates = {}
        for field, raw in patch.items():
            if field == "reason":
                continue
            if not isinstance(raw, str) or len(raw) > 200:
                raise ValueError("Реквизит должен быть строкой до 200 символов")
            value = raw.strip()
            if field in {"code", "name", "manufacturer"} and not value:
                raise ValueError("Обозначение, название и изготовитель не могут быть пустыми")
            if field == "norm" and value:
                value = decimal_text(value)
                if Decimal(value) < 0:
                    raise ValueError("Входимость не может быть отрицательной")
            updates[field] = value
        action = actions[key]
        original = dict(action["position"])
        action["position"].update(updates)
        action["manual_resolution"] = {
            "original": original,
            "patch": updates,
            "reason": patch["reason"].strip(),
        }
        if "manufacturer" in updates:
            action["suppliers"] = [
                {
                    "name": updates["manufacturer"],
                    "contract": action["position"].get("contract", ""),
                    "source_row": action["source_rows"][0],
                }
            ]
        resolved = set()
        if updates.get("manufacturer"):
            resolved.add("MANUFACTURER_REQUIRED")
        if "name" in updates or "code" in updates:
            resolved.add("IDENTITY_CONFLICT")
        if "norm" in updates:
            resolved.add("STRUCTURE_REQUIRED")
        action["errors"] = [e for e in action.get("errors", []) if e["code"] not in resolved]
        for source_key, source in result["sources"].items():
            sheet_index, address = source_key.split(":", 1)
            if (
                int(sheet_index) == action["sheet_index"]
                and coordinate_to_tuple(address)[0] in action["source_rows"]
            ):
                source["position"].update(updates)
                source["manual_resolution"] = action["manual_resolution"]
                if source_key in result["fields"]:
                    result["fields"][source_key]["position"].update(updates)
    return result
